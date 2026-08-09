"""
DynMalTool — Streamlit UI  v5.0
Loads real JSON reports from static_analysis.py / sentry.py / explorer.py / llm_analysis.py.
Runs in DEMO mode when no reports are found — shows the UI structure with placeholder values.

Changes in v5.0:
  - Phase 5 nav tab + THREAT REPORT page (reads phase4_report, all prior reports)
  - APK scan panel moved from sidebar → DASHBOARD (inline, always visible)
  - Report download as Word .docx (python-docx, generated on demand)
  - Screenshots section: dedicated SCREENSHOTS nav tab, reads session screenshots dir
  - Sidebar: APK scan buttons removed; SETTINGS + status only
  - Sentry page: startup_dialogs_handled + packer_chain now displayed
  - Phase 2 runner: --no-explore checkbox + --apk direct-mode text input
  - Phase 3 runner: --apk and --wait inputs
  - Comprehensive error handling: every JSON read, every subprocess, every file op is guarded
  - Graceful degradation everywhere: missing keys never crash the UI
  - _safe_get() helper used throughout to prevent KeyError / TypeError on bad report shapes

Run:  streamlit run app.py
Requires: pip install streamlit python-docx python-dotenv
"""

import streamlit as st
import json
import os
import re
import glob
import subprocess
import threading
import queue as _queue
import sys
import io
import base64
import datetime
from pathlib import Path

# ── optional python-docx import (graceful if absent) ─────────────────────────
try:
    from docx import Document as DocxDocument
    from docx.shared import Pt, RGBColor, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False

# ─── Page config ──────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="DynMal",
    page_icon="🦠",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ══════════════════════════════════════════════════════════════════════════════
#  CSS
# ══════════════════════════════════════════════════════════════════════════════
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Share+Tech+Mono&family=Rajdhani:wght@400;500;600;700&family=Exo+2:wght@300;400;600;700&display=swap');

:root {
    --bg-base:       #080c10;
    --bg-panel:      #0d1117;
    --bg-card:       #111820;
    --bg-card2:      #141c24;
    --border:        #1e2d3d;
    --border-bright: #2a4060;
    --cyan:          #00e5ff;
    --cyan-dim:      #007a8a;
    --magenta:       #ff2d78;
    --magenta-dim:   #7a1038;
    --green:         #00ff88;
    --green-dim:     #005c32;
    --yellow:        #ffc400;
    --red:           #ff4444;
    --orange:        #ff8c00;
    --text-primary:  #e0eaf5;
    --text-secondary:#7a9ab5;
    --text-dim:      #3d5570;
    --font-mono:     'Share Tech Mono', monospace;
    --font-display:  'Rajdhani', sans-serif;
    --font-body:     'Exo 2', sans-serif;
}

html, body, [class*="css"] {
    background-color: var(--bg-base) !important;
    color: var(--text-primary) !important;
    font-family: var(--font-body) !important;
}

#MainMenu, footer, header { visibility: hidden; }
section[data-testid="stSidebar"] > div { padding-top: 0 !important; }

/* ── Sidebar ── */
section[data-testid="stSidebar"] {
    background: var(--bg-panel) !important;
    border-right: 1px solid var(--border) !important;
    min-width: 180px !important;
    max-width: 360px !important;
}
section[data-testid="stSidebar"] > div:last-child {
    background: var(--border-bright) !important;
    width: 3px !important;
    transition: background 0.15s ease !important;
}
section[data-testid="stSidebar"] > div:last-child:hover {
    background: var(--cyan) !important;
    cursor: col-resize !important;
}

.main .block-container {
    padding: 0 1rem !important;
    max-width: 100% !important;
}
.main > div:first-child { padding-top: 0 !important; }

div[data-testid="stSidebar"] .stRadio > div { gap: 0 !important; }
div[data-testid="stSidebar"] .stRadio div[data-baseweb="radio"] > div:first-child { display: none !important; }
div[data-testid="stSidebar"] .stRadio [data-testid="stMarkdownContainer"] { display: none !important; }

div[data-testid="stSidebar"] .stRadio label {
    display: flex !important; align-items: center !important; gap: 12px !important;
    padding: 11px 20px !important; font-family: var(--font-display) !important;
    font-size: 12px !important; font-weight: 600 !important; letter-spacing: 1.8px !important;
    text-transform: uppercase !important; color: var(--text-secondary) !important;
    border-right: 3px solid transparent !important; border-left: none !important;
    cursor: pointer !important; transition: all 0.15s ease !important;
    background: transparent !important; border-radius: 0 !important; width: 100% !important;
}
div[data-testid="stSidebar"] .stRadio label:hover {
    color: var(--text-primary) !important;
    background: rgba(255,255,255,0.03) !important;
}
div[data-testid="stSidebar"] [data-checked="true"] label,
div[data-testid="stSidebar"] .stRadio label[data-baseweb="radio"]:has(input:checked) {
    color: var(--magenta) !important;
    border-right-color: var(--magenta) !important;
    background: rgba(255,45,120,0.06) !important;
}

div[data-testid="stSidebar"] .stRadio [data-testid="stMarkdownContainer"] {
    display: flex !important; align-items: center !important; justify-content: center !important;
    width: 28px !important; min-width: 28px !important; height: 28px !important;
    border: 1px solid var(--border-bright) !important; border-radius: 4px !important;
    background: var(--bg-card) !important; font-size: 13px !important; flex-shrink: 0 !important;
    transition: all 0.15s ease !important;
}
div[data-testid="stSidebar"] .stRadio [data-testid="stMarkdownContainer"] p { margin: 0 !important; line-height: 1 !important; }
div[data-testid="stSidebar"] [data-checked="true"] [data-testid="stMarkdownContainer"],
div[data-testid="stSidebar"] .stRadio label[data-baseweb="radio"]:has(input:checked) [data-testid="stMarkdownContainer"] {
    border-color: var(--magenta) !important;
    background: rgba(255,45,120,0.12) !important;
}
div[data-testid="stSidebar"] .stRadio label:hover [data-testid="stMarkdownContainer"] {
    border-color: var(--border-bright) !important;
    background: var(--bg-card2) !important;
}

.sb-brand { padding: 18px 20px 14px; border-bottom: 1px solid var(--border); margin-bottom: 6px; }
.sb-brand-tag  { font-family: var(--font-mono); font-size: 9px; letter-spacing: 2px; color: var(--magenta); text-transform: uppercase; margin-bottom: 3px; }
.sb-brand-name { font-family: var(--font-display); font-size: 20px; font-weight: 700; letter-spacing: 2px; color: var(--cyan); }
.sb-brand-sub  { font-family: var(--font-mono); font-size: 8px; color: var(--text-dim); letter-spacing: 1px; margin-top: 1px; }

.sb-divider { height: 1px; background: var(--border); margin: 8px 0; }

div[data-testid="stSidebar"] .stButton > button {
    display: flex !important; align-items: center !important; justify-content: center !important;
    gap: 8px !important; width: calc(100% - 32px) !important; margin: 4px 16px !important;
    padding: 9px 0 !important; text-align: center !important;
    font-family: var(--font-display) !important; font-size: 12px !important;
    font-weight: 700 !important; letter-spacing: 2px !important; text-transform: uppercase !important;
    color: var(--text-primary) !important; background: var(--bg-card2) !important;
    border: 1px solid var(--border-bright) !important; border-radius: 3px !important;
    cursor: pointer !important; transition: all 0.15s ease !important;
}
div[data-testid="stSidebar"] .stButton > button:hover {
    border-color: var(--cyan) !important;
    color: var(--cyan) !important;
    background: rgba(0,229,255,0.06) !important;
}

div[data-testid="stSidebar"] [data-testid="stButton"] button[kind="secondary"] {
    justify-content: flex-start !important;
    padding: 10px 20px !important;
    margin: 0 !important;
    width: 100% !important;
    border: none !important;
    border-radius: 0 !important;
    background: transparent !important;
    color: var(--text-dim) !important;
    font-size: 12px !important;
    letter-spacing: 1.8px !important;
}
div[data-testid="stSidebar"] [data-testid="stButton"] button[kind="secondary"]:hover {
    color: var(--text-primary) !important;
    background: rgba(255,255,255,0.03) !important;
    border: none !important;
}

.demo-banner {
    background: rgba(255,196,0,0.08); border: 1px solid rgba(255,196,0,0.3); border-radius: 4px;
    padding: 8px 14px; margin: 0 16px 10px; font-family: var(--font-mono); font-size: 9px;
    color: var(--yellow); letter-spacing: 1px; text-align: center;
}

/* ══ LAYOUT ════════════════════════════════════════════════════════════════ */
.top-bar { background: var(--bg-panel); border-bottom: 1px solid var(--border); padding: 12px 28px; display: flex; align-items: center; justify-content: space-between; margin-bottom: 0; margin-left: -1rem; margin-right: -1rem; }
.brand { font-family: var(--font-display); font-size: 22px; font-weight: 700; letter-spacing: 2px; color: var(--cyan); }
.brand span { color: var(--text-primary); }
.phase-badge { font-family: var(--font-display); font-size: 10px; font-weight: 700; letter-spacing: 2px; text-transform: uppercase; padding: 4px 12px; border-radius: 2px; border: 1px solid; display: inline-block; margin-right: 8px; }
.phase-badge.complete { border-color: var(--green);  color: var(--green);    background: rgba(0,255,136,0.06); }
.phase-badge.pending  { border-color: var(--border); color: var(--text-dim); }

/* ══ PANELS & CARDS ════════════════════════════════════════════════════════ */
.dmt-panel { background: var(--bg-panel); border: 1px solid var(--border); border-radius: 4px; padding: 20px; margin-bottom: 16px; }
.dmt-card  { background: var(--bg-card);  border: 1px solid var(--border); border-radius: 4px; padding: 16px; margin-bottom: 12px; }
.dmt-card-accent-red   { border-left: 3px solid var(--red) !important; }
.dmt-card-accent-yel   { border-left: 3px solid var(--yellow) !important; }
.dmt-card-accent-cyan  { border-left: 3px solid var(--cyan) !important; }
.dmt-card-accent-green { border-left: 3px solid var(--green) !important; }
.dmt-card-accent-mag   { border-left: 3px solid var(--magenta) !important; }

.section-title { font-family: var(--font-display); font-size: 11px; font-weight: 700; letter-spacing: 3px; text-transform: uppercase; color: var(--text-secondary); border-bottom: 1px solid var(--border); padding-bottom: 8px; margin-bottom: 14px; display: flex; align-items: center; gap: 8px; }

.metric-box { background: var(--bg-card); border: 1px solid var(--border); border-radius: 4px; padding: 16px 18px; text-align: center; }
.metric-val { font-family: var(--font-display); font-size: 32px; font-weight: 700; line-height: 1; margin-bottom: 4px; }
.metric-lbl { font-family: var(--font-display); font-size: 9px; font-weight: 700; letter-spacing: 2px; text-transform: uppercase; color: var(--text-dim); }

.sev-badge { font-family: var(--font-display); font-size: 9px; font-weight: 700; letter-spacing: 1.5px; text-transform: uppercase; padding: 2px 8px; border-radius: 2px; display: inline-block; }
.sev-high    { background: rgba(255,68,68,0.15);   color: var(--red); }
.sev-warning { background: rgba(255,196,0,0.15);   color: var(--yellow); }
.sev-info    { background: rgba(0,229,255,0.12);   color: var(--cyan); }
.sev-good    { background: rgba(0,255,136,0.12);   color: var(--green); }

.ioc-type { font-family: var(--font-display); font-size: 9px; font-weight: 700; letter-spacing: 1.5px; padding: 2px 7px; border-radius: 2px; text-transform: uppercase; }
.ioc-type.domain { background: rgba(0,229,255,0.12); color: var(--cyan); }
.ioc-type.ip     { background: rgba(255,196,0,0.12); color: var(--yellow); }
.ioc-type.url    { background: rgba(255,140,0,0.12); color: var(--orange); }
.ioc-type.file   { background: rgba(255,68,68,0.12); color: var(--red); }

.perm-row { display: flex; align-items: center; gap: 10px; margin-bottom: 8px; }
.perm-name { font-family: var(--font-mono); font-size: 11px; color: var(--text-secondary); min-width: 200px; }
.perm-bar-bg { flex: 1; height: 4px; background: var(--border); border-radius: 2px; }
.perm-bar-fill { height: 100%; border-radius: 2px; }
.perm-score { font-family: var(--font-mono); font-size: 11px; min-width: 24px; text-align: right; }

.thread-row { display: flex; align-items: center; gap: 10px; padding: 7px 0; border-bottom: 1px solid var(--border); }
.thread-dot { width: 8px; height: 8px; border-radius: 50%; flex-shrink: 0; }
.thread-dot.active { background: var(--green); box-shadow: 0 0 6px var(--green); animation: pulse 1.5s ease-in-out infinite; }
.thread-dot.idle   { background: var(--text-dim); }
.thread-dot.error  { background: var(--red); }
.thread-name   { font-family: var(--font-mono); font-size: 12px; color: var(--text-secondary); flex: 1; }
.thread-events { font-family: var(--font-mono); font-size: 11px; color: var(--cyan); }

.ioc-table { width: 100%; border-collapse: collapse; font-family: var(--font-mono); font-size: 12px; }
.ioc-table th { font-family: var(--font-display); font-size: 9px; letter-spacing: 2px; text-transform: uppercase; color: var(--text-dim); padding: 8px 10px; text-align: left; border-bottom: 1px solid var(--border); }
.ioc-table td { padding: 9px 10px; border-bottom: 1px solid var(--border); color: var(--text-primary); }
.ioc-table tr:last-child td { border-bottom: none; }

.log-terminal { background: #060a0e; border: 1px solid var(--border); border-radius: 4px; padding: 14px 16px; font-family: var(--font-mono); font-size: 12px; line-height: 1.7; max-height: 520px; overflow-y: auto; }
.log-line { display: flex; gap: 12px; }
.log-ts          { color: var(--text-dim);    min-width: 68px; }
.log-level-info  { color: var(--cyan);        min-width: 50px; font-weight: 700; }
.log-level-warn  { color: var(--yellow);      min-width: 50px; font-weight: 700; }
.log-level-error { color: var(--red);         min-width: 50px; font-weight: 700; }
.log-level-sys   { color: var(--magenta);     min-width: 50px; font-weight: 700; }
.log-level-live  { color: var(--green);       min-width: 50px; font-weight: 700; }
.log-level-adb   { color: var(--orange);      min-width: 50px; font-weight: 700; }
.log-msg         { color: var(--text-primary); }

.state-node { background: var(--bg-card); border: 1px solid var(--border-bright); border-radius: 3px; padding: 4px 10px; font-family: var(--font-mono); font-size: 11px; color: var(--text-secondary); display: inline-block; }
.state-node.visited { border-color: var(--green-dim); color: var(--green); }
.state-node.last    { border-color: var(--cyan); color: var(--cyan); }

.no-data { display: flex; flex-direction: column; align-items: center; justify-content: center; padding: 32px 20px; color: var(--text-dim); font-family: var(--font-mono); font-size: 11px; letter-spacing: 1px; text-align: center; }
.no-data-icon { font-size: 28px; margin-bottom: 10px; opacity: 0.4; }

/* Screenshot grid */
.ss-grid { display: grid; grid-template-columns: repeat(auto-fill, minmax(200px, 1fr)); gap: 12px; padding: 4px 0; }
.ss-card { background: var(--bg-card); border: 1px solid var(--border); border-radius: 4px; overflow: hidden; cursor: pointer; transition: border-color 0.15s; }
.ss-card:hover { border-color: var(--cyan); }
.ss-label { font-family: var(--font-mono); font-size: 9px; color: var(--text-dim); padding: 6px 8px; border-top: 1px solid var(--border); white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }

/* New scan inline panel */
.scan-panel { background: var(--bg-card); border: 1px solid var(--border-bright); border-radius: 4px; padding: 20px; margin-bottom: 16px; }
.scan-panel-title { font-family: var(--font-display); font-size: 11px; font-weight: 700; letter-spacing: 3px; text-transform: uppercase; color: var(--cyan); margin-bottom: 16px; }

div[data-testid="stSidebar"] .streamlit-expanderHeader {
    background: var(--bg-card2) !important;
    border: 1px solid var(--border-bright) !important;
    border-radius: 3px !important;
    font-family: var(--font-display) !important;
    font-size: 12px !important;
    font-weight: 700 !important;
    letter-spacing: 2px !important;
    color: var(--text-secondary) !important;
    text-transform: uppercase !important;
    padding: 9px 14px !important;
    margin: 4px 16px !important;
    width: calc(100% - 32px) !important;
}
div[data-testid="stSidebar"] .streamlit-expanderContent {
    background: var(--bg-panel) !important;
    border: 1px solid var(--border) !important;
    border-top: none !important;
    margin: 0 16px !important;
    padding: 12px !important;
}

.sb-status-block { padding: 10px 20px 8px; }
.sb-status-row { display: flex; align-items: center; gap: 8px; font-family: var(--font-mono); font-size: 9px; color: var(--text-dim); letter-spacing: 1px; padding: 3px 0; }
.sb-dot-g { width: 6px; height: 6px; border-radius: 50%; background: var(--green); box-shadow: 0 0 4px var(--green); flex-shrink: 0; }
.sb-dot-y { width: 6px; height: 6px; border-radius: 50%; background: var(--yellow); flex-shrink: 0; }
.sb-dot-r { width: 6px; height: 6px; border-radius: 50%; background: var(--red); flex-shrink: 0; }
.sb-dot-d { width: 6px; height: 6px; border-radius: 50%; background: var(--text-dim); flex-shrink: 0; }

@keyframes pulse { 0%, 100% { opacity: 1; } 50% { opacity: 0.4; } }
@keyframes scanline { from { transform: translateY(-100%); } to { transform: translateY(100vh); } }
.scanline-overlay { pointer-events: none; position: fixed; top: 0; left: 0; right: 0; height: 2px; background: linear-gradient(transparent, rgba(0,229,255,0.06), transparent); animation: scanline 6s linear infinite; z-index: 9999; }

::-webkit-scrollbar { width: 4px; height: 4px; }
::-webkit-scrollbar-track { background: var(--bg-base); }
::-webkit-scrollbar-thumb { background: var(--border-bright); border-radius: 2px; }
</style>
<div class="scanline-overlay"></div>
""", unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════════
#  DATA LOADING
# ══════════════════════════════════════════════════════════════════════════════

@st.cache_data(ttl=10)
def load_json(path: str):
    """Load a JSON file. Returns None (never raises) on any failure."""
    try:
        p = Path(path)
        if not p.exists():
            return None
        return json.loads(p.read_text(errors="replace"))
    except json.JSONDecodeError as e:
        # Silently degrade — surface as None; page will show 'not run' state
        return None
    except OSError:
        return None
    except Exception:
        return None


def _safe_get(obj, *keys, default=None):
    """
    Safe nested key access on dicts/lists without raising.
    _safe_get(d, "a", "b", "c") == d.get("a",{}).get("b",{}).get("c", default)
    """
    cur = obj
    for k in keys:
        if cur is None:
            return default
        try:
            if isinstance(cur, dict):
                cur = cur.get(k, None)
            elif isinstance(cur, (list, tuple)) and isinstance(k, int):
                cur = cur[k] if k < len(cur) else None
            else:
                return default
        except Exception:
            return default
    return cur if cur is not None else default


def find_all_sessions() -> list:
    """Return all session dirs that contain at least one phase report, newest first."""
    _PHASE_FILES = ("static_report.json", "sentry_report.json", "explorer_report.json", "phase4_report.json")
    try:
        p = Path("sessions")
        if not p.exists():
            return []
        dirs = []
        for d in p.iterdir():
            try:
                if d.is_dir() and any((d / f).exists() for f in _PHASE_FILES):
                    dirs.append(d)
            except OSError:
                continue
        return sorted(dirs, key=lambda d: d.stat().st_mtime, reverse=True)
    except Exception:
        return []


def find_in_session(session_dir, filename: str):
    """Recursively find a file in a session dir. Returns first hit or None."""
    if session_dir is None:
        return None
    try:
        hits = list(session_dir.rglob(filename))
        return hits[0] if hits else None
    except Exception:
        return None


def find_screenshots(session_dir) -> list:
    """
    Return sorted list of PNG paths from the main screenshots dir
    and any dropper screenshot subdirs.
    """
    if session_dir is None:
        return []
    try:
        paths = []
        # Main screenshots dir
        main_ss = session_dir / "screenshots"
        if main_ss.is_dir():
            paths.extend(sorted(main_ss.glob("*.png")))
        # Dropper subdirs: session/droppers/<pkg>/screenshots/*.png
        droppers_root = session_dir / "droppers"
        if droppers_root.is_dir():
            for pkg_dir in sorted(droppers_root.iterdir()):
                try:
                    ss_sub = pkg_dir / "screenshots"
                    if ss_sub.is_dir():
                        paths.extend(sorted(ss_sub.glob("*.png")))
                except Exception:
                    continue
        return paths
    except Exception:
        return []


def session_display_name(session_dir) -> str:
    """Human-readable label: '<stem> · YYYY-MM-DD HH:MM'"""
    try:
        name = session_dir.name
        parts = name.rsplit("_", 2)
        if len(parts) == 3:
            stem, date, tpart = parts
            label_ts = f"{date[:4]}-{date[4:6]}-{date[6:]} {tpart[:2]}:{tpart[2:4]}"
            return f"{stem}  ·  {label_ts}"
        return name
    except Exception:
        return str(session_dir.name) if session_dir else "—"


# ── Session selector ──────────────────────────────────────────────────────────
all_sessions = find_all_sessions()

if "selected_session_name" not in st.session_state:
    st.session_state.selected_session_name = (
        all_sessions[0].name if all_sessions else None
    )

selected_session: Path | None = None
if all_sessions:
    valid_names = {d.name: d for d in all_sessions}
    if st.session_state.selected_session_name not in valid_names:
        st.session_state.selected_session_name = all_sessions[0].name
    selected_session = valid_names.get(st.session_state.selected_session_name)

# ── Load reports for the selected session ────────────────────────────────────
if selected_session:
    static_report   = load_json(str(selected_session / "static_report.json"))
    sentry_report   = load_json(str(selected_session / "sentry_report.json"))
    _ex_path        = find_in_session(selected_session, "explorer_report.json")
    explorer_report = load_json(str(_ex_path)) if _ex_path else None
    phase4_report   = load_json(str(selected_session / "phase4_report.json"))
    phase5_report   = load_json(str(selected_session / "phase5_report.json"))
    _logcat_path    = find_in_session(selected_session, "logcat.txt")
    logcat_path     = str(_logcat_path) if _logcat_path else None
    screenshots     = find_screenshots(selected_session)
else:
    static_report = sentry_report = explorer_report = None
    phase4_report = phase5_report = None
    logcat_path   = None
    screenshots   = []

DEMO_MODE = not any([static_report, sentry_report, explorer_report, phase4_report])

# ── Demo-mode empty skeletons ─────────────────────────────────────────────────
_EMPTY_STATIC = {
    "package":"—","apk_name":"—","apk_size_mb":0,"min_sdk":0,"target_sdk":0,
    "main_activity":"—","has_launcher":False,"pid":None,"all_pids":[],
    "summary":{
        "permission_risk_score":0,"top_risk_permissions":[],"total_indicators":0,
        "manifest_issues":[],"certificate":{"self_signed":False,"expired":False,
        "weak_algorithm":False,"algorithm":"—","fingerprint":"—"},
        "code_patterns":{"threats":[],"defences":[]},
        "indicators":{"urls":[],"ips":[],"domains":[]},"screenshot_tier":3,
    },
    "device":{"adb_connected":False,"frida_running":False,"mitmproxy_cert":False,
              "rooted":False,"screenshot_tier":3,"model":"—","android_version":"—"},
}
_EMPTY_SENTRY = {
    "package_name": "—", "package": "—", "pid": None,
    "session_start": "—", "session_end": "—",
    "all_pids": [], "log_file": "", "duration_seconds": 0,
    "log_events": [], "network_connections": [], "exfiltration_events": [],
    "dropper_events": [], "packer_events": [],
    "root_detected": False, "crashed": False, "crash_timestamp": "",
    "restart_attempts": 0, "startup_dialogs_handled": 0,
    "inotifywait_available": False, "mitm_active": False, "ssl_bypass_active": False,
    "errors": [], "warnings": [], "suggested_commands": [],
    "threads": [{"name": n, "status": "idle", "events": 0} for n in
                ["LogcatStreamer", "RootDetectionMonitor", "FilesystemWatcher",
                 "PackageInstallerWatcher", "NetstatPoller", "MitmproxyRunner",
                 "DropperQueueConsumer", "CrashRecovery"]],
    "summary": {
        "abuse_scores": {}, "top_abused_permissions": [], "exfiltration_types": [],
        "unique_remote_ips": [], "dropper_chain": [], "packer_chain": [],
        "root_detected": False, "crashes": 0, "network_connections": 0,
        "total_log_events": 0, "logcat_events": 0, "startup_dialogs_handled": 0,
        "total_network_connections": 0, "total_exfiltration_events": 0,
        "total_dropper_events": 0, "total_packer_events": 0,
        "inotifywait_used": False, "mitm_active": False, "ssl_bypass_active": False,
        "permission_hit_counts": {}, "exfiltration_types_list": [],
    },
}
_EMPTY_EXPLORER = {
    "package_name": "—", "package": "—", "session_dir": "—",
    "states_visited": 0, "total_actions": 0,
    "permission_dialogs_accepted": 0, "forms_filled": 0, "crashes_recovered": 0,
    "form_submits_attempted": 0, "form_submits_succeeded": 0,
    "screenshot_tier": 3, "states": [], "actions": [],
    "form_submit_failures": [], "errors": [], "warnings": [],
    "summary": {
        "states_visited": 0, "total_actions": 0,
        "deepest_depth": 0, "deepest_state": 0,
        "permission_dialogs_accepted": 0, "forms_filled": 0,
        "crash_recoveries": 0, "crashes_recovered": 0,
        "screenshot_tier": 3, "unique_activities": [],
        "permissions_requested": [],
        "form_submits_attempted": 0, "form_submits_succeeded": 0,
        "form_submit_failure_count": 0, "form_submit_failures": [],
        "human_interventions": 0, "state_hashes": [],
    },
}
_EMPTY_PHASE4 = {
    "llm_status": "not_run",
    "network_iocs": [],
    "suspicious_iocs": [],
    "dropper_chain": [],
    "packer_chain": [],
    # flat aliases kept for backward compat — all real data lives in llm_verdict
    "threat_class": "—", "confidence": 0, "threat_score": 0,
    "kill_chain_stage": "—", "analyst_notes": "—",
    "evidence": [], "iocs": [], "data_collected": [], "recommended_actions": [],
    "llm_verdict": {
        "threat_class": "—",
        "confidence": "LOW",
        "threat_score": 0,
        "kill_chain_stage": "—",
        "evidence": [],
        "iocs": {"domains": [], "ips": [], "urls": [], "permissions": [], "code_patterns": []},
        "data_collected": [],
        "recommended_actions": [],
        "analyst_notes": "—",
    },
}
_EMPTY_PHASE5 = {
    "status":"not_run","generated_at":"—",
    "executive_summary":"—","threat_classification":"—","kill_chain_analysis":"—",
    "network_iocs":[],"permission_abuse":[],"dropper_packer_chain":[],
    "ui_behaviour":"—","recommendations":[],
}

S  = static_report   or _EMPTY_STATIC
SE = sentry_report   or _EMPTY_SENTRY
EX = explorer_report or _EMPTY_EXPLORER
P4 = phase4_report   or _EMPTY_PHASE4
P5 = phase5_report   or _EMPTY_PHASE5


# ══════════════════════════════════════════════════════════════════════════════
#  HELPERS
# ══════════════════════════════════════════════════════════════════════════════

def section_title(icon, text):
    st.markdown(f'<div class="section-title"><span>{icon}</span> {text}</div>', unsafe_allow_html=True)

def sev_badge(sev: str) -> str:
    sev = (sev or "info").lower()
    cls = {"high":"sev-high","warning":"sev-warning","info":"sev-info","good":"sev-good"}.get(sev,"sev-info")
    return f'<span class="sev-badge {cls}">{sev.upper()}</span>'

def no_data(msg="No data — run the pipeline first"):
    st.markdown(f'<div class="no-data"><div class="no-data-icon">⬡</div>{msg}</div>', unsafe_allow_html=True)

def kv(label, value, color="var(--text-primary)"):
    safe_val = str(value) if value is not None else "—"
    return (f'<div style="display:flex;align-items:center;justify-content:space-between;'
            f'padding:7px 0;border-bottom:1px solid var(--border);">'
            f'<span style="font-family:var(--font-display);font-size:12px;color:var(--text-secondary);">{label}</span>'
            f'<span style="font-family:var(--font-mono);font-size:12px;color:{color};">{safe_val}</span></div>')

def phase_badge(label, report):
    if report:
        return f'<span class="phase-badge complete">{label} ✓</span>'
    return f'<span class="phase-badge pending">{label} —</span>'

def dot(ok):
    return '<div class="sb-dot-g"></div>' if ok else '<div class="sb-dot-r"></div>'

def _escape_html(s: str) -> str:
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

def render_log_lines(lines):
    html = '<div class="log-terminal">'
    lvl_map = {"INFO":"log-level-info","WARN":"log-level-warn","ERROR":"log-level-error",
               "SYSTEM":"log-level-sys","LIVE":"log-level-live","ADB":"log-level-adb"}
    for ts, level, msg in lines:
        cls = lvl_map.get(level, "log-level-info")
        safe_msg = _escape_html(msg)
        html += (f'<div class="log-line">'
                 f'<span class="log-ts">[{ts}]</span>'
                 f'<span class="{cls}">{level}</span>'
                 f'<span class="log-msg">{safe_msg}</span></div>')
    html += '</div>'
    return html


# ══════════════════════════════════════════════════════════════════════════════
#  SESSION STATE
# ══════════════════════════════════════════════════════════════════════════════
for _k, _v in [
    ("show_settings", False),
    ("phase_proc",    None),
    ("phase_output",  []),
    ("phase_label",   ""),
    ("phase_done",    False),
    ("phase_error",   False),
    ("phase_out_q",   None),
]:
    if _k not in st.session_state:
        st.session_state[_k] = _v


def _start_phase(cmd: list, label: str) -> None:
    """Launch a phase subprocess and wire up a reader thread."""
    try:
        proc = subprocess.Popen(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            bufsize=1,
        )
    except FileNotFoundError as e:
        st.error(f"Could not launch {label}: {e}. Check that all phase scripts are in the working directory.", icon="🚨")
        return
    except Exception as e:
        st.error(f"Failed to start {label}: {e}", icon="🚨")
        return

    q: _queue.Queue = _queue.Queue()

    def _reader():
        try:
            for line in proc.stdout:
                q.put(line.rstrip())
        except Exception:
            pass
        finally:
            try:
                proc.wait()
            except Exception:
                pass
            q.put(None)  # sentinel

    t = threading.Thread(target=_reader, daemon=True)
    t.start()

    st.session_state.phase_proc   = proc
    st.session_state.phase_out_q  = q
    st.session_state.phase_output = []
    st.session_state.phase_label  = label
    st.session_state.phase_done   = False
    st.session_state.phase_error  = False


def _poll_phase() -> bool:
    """Drain the queue into phase_output. Returns True once the process finishes."""
    q = st.session_state.phase_out_q
    if q is None:
        return True
    finished = False
    while True:
        try:
            line = q.get_nowait()
        except _queue.Empty:
            break
        if line is None:
            finished = True
            proc = st.session_state.phase_proc
            if proc is not None:
                try:
                    st.session_state.phase_error = (proc.returncode != 0)
                except Exception:
                    pass
            break
        st.session_state.phase_output.append(line)
    return finished


def _phase_running() -> bool:
    return (
        st.session_state.phase_proc is not None
        and not st.session_state.phase_done
    )


def _render_phase_runner(next_label: str, next_cmd: list) -> None:
    """Render the 'Run Next Phase' button + live output panel."""
    running = _phase_running()
    done    = st.session_state.phase_done

    btn_col, _ = st.columns([1, 4])
    with btn_col:
        if st.button(
            f"▶  RUN {next_label.upper()}",
            key=f"run_{next_label.replace(' ','_').lower()}_{id(next_cmd)}",
            disabled=running,
            use_container_width=True,
        ):
            _start_phase(next_cmd, next_label)
            st.rerun()

    if running or (done and st.session_state.phase_output):
        finished = _poll_phase()
        if finished and not st.session_state.phase_done:
            st.session_state.phase_done = True
            st.cache_data.clear()
            st.session_state.selected_session_name = None

        label_disp = st.session_state.phase_label
        lines      = st.session_state.phase_output[-120:]

        status_color = (
            "var(--red)"   if st.session_state.phase_error else
            "var(--green)" if st.session_state.phase_done  else
            "var(--cyan)"
        )
        status_text = (
            "FAILED"    if st.session_state.phase_error else
            "COMPLETE"  if st.session_state.phase_done  else
            "RUNNING…"
        )
        lines_html = "".join(
            '<div style="font-family:var(--font-mono);font-size:11px;line-height:1.6;">'
            + _escape_html(l) + "</div>" for l in lines
        )
        st.markdown(f"""
        <div style="background:#060a0e;border:1px solid var(--border);border-radius:4px;
                    padding:12px 16px;margin:8px 0 4px;max-height:300px;overflow-y:auto;">
          <div style="display:flex;align-items:center;justify-content:space-between;
                      margin-bottom:8px;border-bottom:1px solid var(--border);padding-bottom:6px;">
            <span style="font-family:var(--font-display);font-size:10px;letter-spacing:2px;
                         color:var(--text-dim);">{label_disp} OUTPUT</span>
            <span style="font-family:var(--font-mono);font-size:10px;color:{status_color};">{status_text}</span>
          </div>
          <div>{lines_html}</div>
        </div>""", unsafe_allow_html=True)

        if running:
            st.rerun()


# ══════════════════════════════════════════════════════════════════════════════
#  WORD DOCUMENT REPORT GENERATOR
# ══════════════════════════════════════════════════════════════════════════════

def _docx_heading(doc, text: str, level: int = 1):
    """Add a heading paragraph with our dark-theme compatible colours."""
    p = doc.add_heading(text, level=level)
    run = p.runs[0] if p.runs else p.add_run(text)
    run.font.color.rgb = RGBColor(0x00, 0xE5, 0xFF) if level == 1 else RGBColor(0xFF, 0x2D, 0x78)
    return p


def _docx_kv_table(doc, rows: list, col_widths=(4.5, 4.5)):
    """Add a two-column key-value table."""
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for i, (key, val) in enumerate(rows):
        row = table.rows[i]
        row.cells[0].text = str(key)
        row.cells[1].text = str(val) if val is not None else "—"
        for j, w in enumerate(col_widths):
            row.cells[j].width = Inches(w)
    return table


def build_word_report(
    static_r, sentry_r, explorer_r, phase4_r, session_dir
) -> bytes | None:
    """
    Generate a .docx threat report from all available phase reports.
    Returns the docx bytes, or None if python-docx is unavailable.
    """
    if not _DOCX_AVAILABLE:
        return None

    doc = DocxDocument()

    # --- Title page ---
    title = doc.add_heading("DynMalTool — Threat Analysis Report", 0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x00, 0xE5, 0xFF)

    pkg = _safe_get(static_r, "package", default="Unknown Package")
    ts  = _safe_get(static_r, "timestamp", default=datetime.datetime.now().strftime("%Y-%m-%d %H:%M"))
    sess = str(session_dir.name) if session_dir else "—"
    doc.add_paragraph(f"Package: {pkg}")
    doc.add_paragraph(f"Session: {sess}")
    doc.add_paragraph(f"Generated: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_page_break()

    # --- Phase 4 Executive Summary ---
    _docx_heading(doc, "1. Executive Summary")
    if phase4_r:
        _v4 = _safe_get(phase4_r, "llm_verdict", default={}) or {}
        threat_class = _safe_get(_v4, "threat_class", default="—") or _safe_get(phase4_r, "threat_class", default="—")
        threat_score = _safe_get(_v4, "threat_score", default=0) or _safe_get(phase4_r, "threat_score", default=0)
        _conf_raw    = _safe_get(_v4, "confidence", default="") or _safe_get(phase4_r, "confidence", default=0)
        if isinstance(_conf_raw, str):
            confidence = {"LOW": 25, "MEDIUM": 55, "HIGH": 85, "CRITICAL": 95}.get(_conf_raw.upper(), 0)
        else:
            try:
                confidence = int(_conf_raw)
            except (TypeError, ValueError):
                confidence = 0
        kill_chain   = _safe_get(_v4, "kill_chain_stage", default="—") or _safe_get(phase4_r, "kill_chain_stage", default="—")
        _docx_kv_table(doc, [
            ("Threat Classification", threat_class),
            ("Threat Score",          threat_score),
            ("Confidence",            f"{confidence}%"),
            ("Kill Chain Stage",      kill_chain),
            ("LLM Status",            _safe_get(phase4_r, "llm_status", default="—")),
        ])
        doc.add_paragraph()
        analyst_notes = _safe_get(_v4, "analyst_notes", default="") or _safe_get(phase4_r, "analyst_notes", default="")
        if analyst_notes and str(analyst_notes) not in ("—", ""):
            _docx_heading(doc, "Analyst Notes", level=2)
            doc.add_paragraph(str(analyst_notes))
    else:
        doc.add_paragraph("Phase 4 (LLM Analysis) not yet run.")
    doc.add_page_break()

    # --- Static Analysis ---
    _docx_heading(doc, "2. Static Analysis (Phase 1)")
    if static_r:
        summ = _safe_get(static_r, "summary", default={})
        _docx_kv_table(doc, [
            ("APK File",        _safe_get(static_r, "apk_name", default="—")),
            ("Size (MB)",       _safe_get(static_r, "apk_size_mb", default="—")),
            ("Min SDK",         _safe_get(static_r, "min_sdk", default="—")),
            ("Target SDK",      _safe_get(static_r, "target_sdk", default="—")),
            ("Main Activity",   _safe_get(static_r, "main_activity", default="—")),
            ("Permission Risk", _safe_get(summ, "permission_risk_score", default=0)),
        ])
        doc.add_paragraph()
        top_perms = _safe_get(summ, "top_risk_permissions", default=[])
        if top_perms:
            _docx_heading(doc, "Dangerous Permissions", level=2)
            for p in top_perms:
                if isinstance(p, dict):
                    doc.add_paragraph(
                        f"  {p.get('name','—')}  (risk: {p.get('risk_score','?')})",
                        style="List Bullet"
                    )
        manifest_issues = _safe_get(summ, "manifest_issues", default=[])
        if manifest_issues:
            _docx_heading(doc, "Manifest Security Issues", level=2)
            for issue in manifest_issues:
                if isinstance(issue, dict):
                    sev = issue.get("severity","info").upper()
                    title_i = issue.get("title","—")
                    desc_i  = issue.get("description","")
                    doc.add_paragraph(f"[{sev}] {title_i}: {desc_i}", style="List Bullet")
    else:
        doc.add_paragraph("Phase 1 (Static Analysis) not yet run.")
    doc.add_page_break()

    # --- Dynamic Sentry ---
    _docx_heading(doc, "3. Dynamic Sentry (Phase 2)")
    if sentry_r:
        se_sum = _safe_get(sentry_r, "summary", default={})

        def _chain_str(chain):
            """Safely render a dropper/packer chain list to a human-readable string."""
            if not chain:
                return "—"
            parts = []
            for item in chain:
                if isinstance(item, dict):
                    parts.append(item.get("pkg_name") or item.get("package_name") or str(item))
                else:
                    parts.append(str(item))
            return " → ".join(parts)

        # duration from session timestamps when field absent
        _dur = _safe_get(sentry_r, "duration_seconds", default=0)
        if not _dur:
            try:
                from datetime import datetime as _dt
                _s = _safe_get(sentry_r, "session_start", default="")
                _e = _safe_get(sentry_r, "session_end", default="")
                if _s and _e:
                    _dur = round((_dt.fromisoformat(_e) - _dt.fromisoformat(_s)).total_seconds())
            except Exception:
                pass

        _logcat_ev = _safe_get(se_sum, "total_log_events", default=0) or _safe_get(se_sum, "logcat_events", default=0) or len(_safe_get(sentry_r, "log_events", default=[]))
        _net_c     = _safe_get(se_sum, "total_network_connections", default=0) or _safe_get(se_sum, "network_connections", default=0) or len(_safe_get(sentry_r, "network_connections", default=[]))
        _exfil_types = _safe_get(se_sum, "exfiltration_types", default=[])
        if isinstance(_exfil_types, list):
            _exfil_str = ", ".join(str(x) for x in _exfil_types) or "—"
        else:
            _exfil_str = str(_exfil_types) or "—"

        _docx_kv_table(doc, [
            ("Session Start",         _safe_get(sentry_r, "session_start", default="—")),
            ("Session End",           _safe_get(sentry_r, "session_end", default="—")),
            ("Duration (s)",          str(_dur)),
            ("Logcat Events",         str(_logcat_ev)),
            ("Network Connections",   str(_net_c)),
            ("Exfiltration Types",    _exfil_str),
            ("Dropper Chain",         _chain_str(_safe_get(se_sum, "dropper_chain", default=[]))),
            ("Packer Chain",          _chain_str(_safe_get(se_sum, "packer_chain", default=[]))),
            ("Root Detection Fired",  str(_safe_get(se_sum, "root_detected", default=False) or _safe_get(sentry_r, "root_detected", default=False))),
            ("Crashes Detected",      str(_safe_get(se_sum, "crashes", default=0))),
            ("Startup Dialogs",       str(_safe_get(se_sum, "startup_dialogs_handled", default=0) or _safe_get(sentry_r, "startup_dialogs_handled", default=0))),
            ("MITM Active",           str(_safe_get(sentry_r, "mitm_active", default=False))),
            ("SSL Bypass",            str(_safe_get(sentry_r, "ssl_bypass_active", default=False))),
        ])
        doc.add_paragraph()
        net_list = _safe_get(sentry_r, "network_connections", default=[])
        if net_list:
            _docx_heading(doc, "Network Connections", level=2)
            table = doc.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            for i, h in enumerate(["IP / DOMAIN", "PORT", "PROTO", "LOCAL PORT", "FLAG"]):
                hdr[i].text = h
            for conn in net_list[:100]:
                if isinstance(conn, dict):
                    row = table.add_row().cells
                    ip_v  = conn.get("remote_ip") or conn.get("ip") or "—"
                    dom_v = conn.get("domain","")
                    row[0].text = f"{ip_v}" + (f" ({dom_v})" if dom_v else "")
                    row[1].text = str(conn.get("remote_port") or conn.get("port") or "—")
                    row[2].text = str(conn.get("protocol") or conn.get("proto") or "—")
                    row[3].text = str(conn.get("local_port","—"))
                    row[4].text = str(conn.get("flag","—"))

        # Packer events detail
        packer_evts = _safe_get(sentry_r, "packer_events", default=[])
        if packer_evts:
            _docx_heading(doc, "Packer Events", level=2)
            for pe in packer_evts:
                if isinstance(pe, dict):
                    pkg  = pe.get("package_name","—")
                    sha  = pe.get("sha256","—")
                    meth = pe.get("pull_method","—")
                    ts_e = pe.get("timestamp","—")
                    doc.add_paragraph(f"[{ts_e}] {pkg}  SHA256:{sha}  method:{meth}", style="List Bullet")

        # Suggested commands
        cmds = _safe_get(sentry_r, "suggested_commands", default=[])
        if cmds:
            _docx_heading(doc, "Suggested Follow-Up Commands", level=2)
            for cmd in cmds:
                doc.add_paragraph(str(cmd), style="List Bullet")
    else:
        doc.add_paragraph("Phase 2 (Dynamic Sentry) not yet run.")
    doc.add_page_break()

    # --- Explorer ---
    _docx_heading(doc, "4. UI Explorer (Phase 3)")
    if explorer_r:
        ex_sum = _safe_get(explorer_r, "summary", default={})
        _docx_kv_table(doc, [
            ("States Visited",       str(_safe_get(ex_sum, "states_visited",           default=0) or _safe_get(explorer_r, "states_visited", default=0))),
            ("Total Actions",        str(_safe_get(ex_sum, "total_actions",            default=0) or _safe_get(explorer_r, "total_actions", default=0))),
            ("Deepest Depth",        str(_safe_get(ex_sum, "deepest_state",            default=0) or _safe_get(ex_sum, "deepest_depth", default=0))),
            ("Dialogs Accepted",     str(_safe_get(ex_sum, "permission_dialogs_accepted", default=0) or _safe_get(explorer_r, "permission_dialogs_accepted", default=0))),
            ("Forms Filled",         str(_safe_get(ex_sum, "forms_filled",             default=0) or _safe_get(explorer_r, "forms_filled", default=0))),
            ("Form Submits OK",      str(_safe_get(ex_sum, "form_submits_succeeded",   default=0) or _safe_get(explorer_r, "form_submits_succeeded", default=0))),
            ("Crash Recoveries",     str(_safe_get(ex_sum, "crashes_recovered",        default=0) or _safe_get(ex_sum, "crash_recoveries", default=0) or _safe_get(explorer_r, "crashes_recovered", default=0))),
            ("Screenshot Tier",      str(_safe_get(ex_sum, "screenshot_tier",          default=3) or _safe_get(explorer_r, "screenshot_tier", default=3))),
        ])
        activities = _safe_get(ex_sum, "unique_activities", default=[])
        if activities:
            _docx_heading(doc, "Unique Activities Reached", level=2)
            for act in activities:
                doc.add_paragraph(str(act), style="List Bullet")
        fail_list = _safe_get(ex_sum, "form_submit_failures", default=[]) or _safe_get(explorer_r, "form_submit_failures", default=[])
        if fail_list:
            _docx_heading(doc, "Form Submit Failures", level=2)
            for fl in fail_list[:20]:
                if isinstance(fl, dict):
                    btn = fl.get("button","—")
                    ts_f = fl.get("timestamp","—")
                    err_nodes = fl.get("error_nodes",[])
                    err_txt = "; ".join(e.get("text","") for e in err_nodes if isinstance(e, dict)) if err_nodes else ""
                    doc.add_paragraph(f"[{ts_f}] Button: {btn}" + (f" — Error: {err_txt}" if err_txt else ""), style="List Bullet")
    else:
        doc.add_paragraph("Phase 3 (UI Explorer) not yet run.")
    doc.add_page_break()

    # --- IOCs ---
    _docx_heading(doc, "5. Network IOCs (Phase 4)")
    if phase4_r:
        _v4b = _safe_get(phase4_r, "llm_verdict", default={}) or {}
        # flatten iocs from llm_verdict.iocs (dict of lists) + network_iocs (list of dicts)
        _iocs_all = []
        _iocs_dict = _safe_get(_v4b, "iocs", default={}) or {}
        if isinstance(_iocs_dict, dict):
            for ioc_type, ioc_list in _iocs_dict.items():
                if isinstance(ioc_list, list):
                    for v in ioc_list:
                        _iocs_all.append({"type": ioc_type, "value": str(v), "flag": ""})
        elif isinstance(_iocs_dict, list):
            _iocs_all = _iocs_dict
        for ni in _safe_get(phase4_r, "network_iocs", default=[]):
            if isinstance(ni, dict):
                _iocs_all.append({"type": ni.get("ioc_type","ip"), "value": ni.get("value","—"), "flag": ni.get("classification","")})
        if _iocs_all:
            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            for i, h in enumerate(["TYPE", "VALUE", "FLAG"]):
                hdr[i].text = h
            for ioc in _iocs_all:
                row = table.add_row().cells
                if isinstance(ioc, dict):
                    row[0].text = str(ioc.get("type","ioc")).upper()
                    row[1].text = str(ioc.get("value","—"))
                    row[2].text = str(ioc.get("flag",""))
                else:
                    row[0].text = "IOC"
                    row[1].text = str(ioc)
                    row[2].text = ""
        else:
            doc.add_paragraph("No IOCs extracted.")
        recommended = _safe_get(_v4b, "recommended_actions", default=[]) or _safe_get(phase4_r, "recommended_actions", default=[])
        if recommended:
            _docx_heading(doc, "Recommended Actions", level=2)
            for i, action in enumerate(recommended, 1):
                doc.add_paragraph(f"{i}. {action}", style="List Number")
        evidence = _safe_get(_v4b, "evidence", default=[]) or _safe_get(phase4_r, "evidence", default=[])
        if evidence:
            _docx_heading(doc, "LLM Evidence", level=2)
            for ev in evidence:
                if isinstance(ev, dict):
                    doc.add_paragraph(f"[{ev.get('severity','info').upper()}] {ev.get('description', str(ev))}", style="List Bullet")
                else:
                    doc.add_paragraph(str(ev), style="List Bullet")
    else:
        doc.add_paragraph("Phase 4 (LLM Analysis) not yet run.")

    # --- Errors / warnings across all phases ---
    all_errors = []
    for label, report in [("Phase 1", static_r), ("Phase 2", sentry_r), ("Phase 3", explorer_r)]:
        if report:
            errs = _safe_get(report, "errors", default=[])
            if errs:
                for e in errs:
                    all_errors.append(f"[{label}] {e}")
    if all_errors:
        doc.add_page_break()
        _docx_heading(doc, "6. Pipeline Errors & Warnings")
        for err in all_errors:
            doc.add_paragraph(str(err), style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
#  SIDEBAR
# ══════════════════════════════════════════════════════════════════════════════

try:
    device  = S.get("device", {}) or {}
except Exception:
    device  = {}

adb_ok   = bool(_safe_get(device, "adb_connected", default=False))
frida_ok = bool(_safe_get(device, "frida_running", default=False))
mitm_ok  = bool(_safe_get(device, "mitmproxy_cert", default=False) or _safe_get(SE, "mitm_active", default=False))
rooted   = bool(_safe_get(device, "rooted", default=False))
model    = _safe_get(device, "model", default="—")
android  = _safe_get(device, "android_version", default="—")
tier     = _safe_get(device, "screenshot_tier", default=3)
# PID lives in sentry report (Phase 1 no longer owns launch)
_all_pids = _safe_get(SE, "all_pids", default=[]) or _safe_get(S, "all_pids", default=[])
pid_val  = _all_pids[0] if _all_pids else (S.get("pid") if isinstance(S, dict) else None)

with st.sidebar:
    st.markdown("""
    <div class="sb-brand">
        <div class="sb-brand-tag">Cyber_Pipeline</div>
        <div class="sb-brand-name">DynMal</div>
        <div class="sb-brand-sub">VS.0_NEON_CORE · v5.0</div>
    </div>
    """, unsafe_allow_html=True)

    if DEMO_MODE:
        st.markdown('<div class="demo-banner">⚠ DEMO MODE — no reports loaded</div>', unsafe_allow_html=True)

    # ── SESSION SELECTOR ──────────────────────────────────────────────────────
    if all_sessions:
        session_names  = [d.name for d in all_sessions]
        session_labels = [session_display_name(d) for d in all_sessions]
        label_to_name  = dict(zip(session_labels, session_names))

        current_label = session_display_name(selected_session) if selected_session else session_labels[0]
        try:
            _idx = session_labels.index(current_label)
        except ValueError:
            _idx = 0

        chosen_label = st.selectbox(
            "Active session",
            options=session_labels,
            index=_idx,
            label_visibility="collapsed",
            key="session_selector",
        )
        chosen_name = label_to_name.get(chosen_label, session_names[0])
        if chosen_name != st.session_state.selected_session_name:
            st.session_state.selected_session_name = chosen_name
            st.cache_data.clear()
            st.rerun()
    else:
        st.markdown(
            '<div style="font-family:var(--font-mono);font-size:9px;color:var(--text-dim);'
            'padding:6px 20px;">No sessions found — use New Scan below</div>',
            unsafe_allow_html=True,
        )

    st.markdown('<div class="sb-divider" style="margin:4px 0 2px;"></div>', unsafe_allow_html=True)

    # ── NAVIGATION ────────────────────────────────────────────────────────────
    page = st.radio(
        "nav",
        ["📊 ||DASHBOARD",
         "🔬 ||STATIC ANALYSIS",
         "🛡 ||DYNAMIC SENTRY",
         "🧭 ||EXPLORER",
         "🤖 ||LLM ANALYSIS",
         "📄 ||THREAT REPORT",
         "🖼 ||SCREENSHOTS",
         "📋 ||LIVE LOGS"],
        label_visibility="collapsed",
        format_func=lambda x: x.split("||")[1],
    )

    st.markdown('<div class="sb-divider" style="margin:12px 0 6px;"></div>', unsafe_allow_html=True)

    # ── SETTINGS button ───────────────────────────────────────────────────────
    if st.button("⚙  SETTINGS", key="btn_settings", use_container_width=True):
        st.session_state.show_settings = not st.session_state.show_settings

    if st.session_state.show_settings:
        with st.expander("", expanded=True):
            st.markdown('<div style="font-family:var(--font-display);font-size:9px;letter-spacing:2px;color:var(--text-dim);margin-bottom:6px;">ADB SERIAL</div>', unsafe_allow_html=True)
            adb_serial = st.text_input("ADB Serial", value=os.environ.get("ADB_SERIAL",""), placeholder="e.g. 47251VDJH01412", label_visibility="collapsed", key="cfg_adb_serial")
            st.markdown('<div style="font-family:var(--font-display);font-size:9px;letter-spacing:2px;color:var(--text-dim);margin:8px 0 6px;">ANALYSIS DURATION (s)</div>', unsafe_allow_html=True)
            duration_cfg = st.number_input("Analysis duration", value=180, min_value=30, max_value=3600, label_visibility="collapsed", key="cfg_duration")
            st.markdown('<div style="font-family:var(--font-display);font-size:9px;letter-spacing:2px;color:var(--text-dim);margin:8px 0 6px;">MAX EXPLORER STATES</div>', unsafe_allow_html=True)
            max_states_cfg = st.number_input("Max explorer states", value=80, min_value=10, max_value=500, label_visibility="collapsed", key="cfg_states")
            st.markdown('<div style="font-family:var(--font-display);font-size:9px;letter-spacing:2px;color:var(--text-dim);margin:8px 0 6px;">MITM PORT</div>', unsafe_allow_html=True)
            mitm_port_cfg = st.number_input("MITM port", value=8080, min_value=1024, max_value=65535, label_visibility="collapsed", key="cfg_mitm")
            if st.button("Save to .env", key="save_env", use_container_width=True):
                try:
                    existing = {}
                    if Path(".env").exists():
                        for line in Path(".env").read_text(errors="replace").splitlines():
                            if "=" in line and not line.startswith("#"):
                                k, v = line.split("=", 1)
                                existing[k.strip()] = v.strip()
                    existing["ADB_SERIAL"]         = str(adb_serial)
                    existing["MITM_PORT"]           = str(mitm_port_cfg)
                    existing["EXPLORER_MAX_STATES"] = str(max_states_cfg)
                    existing["ANALYSIS_DURATION"]   = str(duration_cfg)
                    Path(".env").write_text("\n".join(f"{k}={v}" for k, v in existing.items()))
                    st.success("Saved!", icon="✓")
                except Exception as e:
                    st.error(f"Could not save .env: {e}", icon="✗")

    st.markdown('<div class="sb-divider" style="margin:6px 0 4px;"></div>', unsafe_allow_html=True)

    # ── Device status ─────────────────────────────────────────────────────────
    pid_row       = f'<div class="sb-status-row"><div class="sb-dot-g"></div>PID · {pid_val}</div>' if pid_val else ""
    sess_name_row = ""
    if selected_session:
        try:
            parts = selected_session.name.rsplit("_", 2)
            if len(parts) == 3:
                d, t = parts[1], parts[2]
                sess_short = f"{d[:4]}-{d[4:6]}-{d[6:]} {t[:2]}:{t[2:4]}"
            else:
                sess_short = selected_session.name
        except Exception:
            sess_short = "—"
        sess_name_row = f'<div class="sb-status-row"><div class="sb-dot-y"></div>SESSION · {sess_short}</div>'

    ss_count_row = f'<div class="sb-status-row"><div class="sb-dot-y"></div>SCREENSHOTS · {len(screenshots)}</div>'

    st.markdown(f"""
    <div class="sb-status-block">
        {sess_name_row}
        <div class="sb-status-row">{dot(adb_ok)} ADB {"CONNECTED" if adb_ok else "DISCONNECTED"}</div>
        <div class="sb-status-row">{dot(frida_ok)} FRIDA {"RUNNING" if frida_ok else "NOT RUNNING"}</div>
        <div class="sb-status-row">{dot(mitm_ok)} MITM {"ACTIVE" if mitm_ok else "NOT DETECTED"}</div>
        <div class="sb-status-row">{dot(rooted)} {"ROOTED" if rooted else "NOT ROOTED"}</div>
        <div class="sb-status-row" style="margin-top:8px;"><div class="sb-dot-y"></div>SCREENSHOT TIER {tier}</div>
        <div class="sb-status-row"><div class="sb-dot-y"></div>DEVICE · {model}</div>
        <div class="sb-status-row"><div class="sb-dot-y"></div>ANDROID · {android}</div>
        {ss_count_row}
        {pid_row}
    </div>
    """, unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════════
#  DASHBOARD
# ══════════════════════════════════════════════════════════════════════════════

if "DASHBOARD" in page:
    summ  = _safe_get(S,  "summary", default={})
    s_sum = _safe_get(SE, "summary", default={})
    e_sum = _safe_get(EX, "summary", default={})

    perm_score = _safe_get(summ, "permission_risk_score", default=0)
    total_ind  = _safe_get(summ, "total_indicators", default=0)
    net_conns  = _safe_get(s_sum, "network_connections", default=0)
    exfil_list = _safe_get(s_sum, "exfiltration_types", default=[])
    states_vis = _safe_get(e_sum, "states_visited", default=0)

    # Phase 4 threat score for top bar — real data is nested under llm_verdict
    ts_val = _safe_get(P4, "llm_verdict", "threat_score", default=0) or _safe_get(P4, "threat_score", default=0)
    try:
        ts_val = int(ts_val)
    except (TypeError, ValueError):
        ts_val = 0
    ts_color = "var(--red)" if ts_val >= 75 else "var(--orange)" if ts_val >= 50 else "var(--yellow)" if ts_val >= 25 else "var(--text-dim)"

    st.markdown(f"""
    <div class="top-bar">
      <div style="display:flex;align-items:center;gap:20px;">
        <span class="brand">Life<span>Cycle</span>View</span>
        {phase_badge("P1", static_report)}
        {phase_badge("P2", sentry_report)}
        {phase_badge("P3", explorer_report)}
        {phase_badge("P4", phase4_report)}
        {phase_badge("P5", phase5_report)}
      </div>
      <div style="display:flex;align-items:center;gap:14px;">
        <span style="font-family:var(--font-mono);font-size:11px;color:var(--text-dim);">{_safe_get(S, "package_name", default="") or _safe_get(S, "package", default="—")}</span>
        {'<span style="font-family:var(--font-display);font-size:18px;font-weight:700;color:' + ts_color + ';">THREAT · ' + str(ts_val) + '</span>' if phase4_report else ""}
        <div style="background:rgba(255,45,120,0.12);border:1px solid var(--magenta);border-radius:3px;padding:5px 16px;font-family:var(--font-display);font-size:11px;font-weight:700;letter-spacing:1.5px;color:var(--magenta);">v5.0</div>
      </div>
    </div>""", unsafe_allow_html=True)

    # ── NEW SCAN panel (inline in dashboard) ──────────────────────────────────
    st.markdown('<div style="padding:20px 28px 0;">', unsafe_allow_html=True)
    with st.expander("⊕  NEW SCAN — Phase 1", expanded=not static_report):
        st.markdown('<div class="scan-panel-title">START NEW ANALYSIS</div>', unsafe_allow_html=True)
        apk_input = st.text_input(
            "APK path",
            placeholder="apks/target.apk  or  /absolute/path/to/sample.apk",
            label_visibility="collapsed",
            key="apk_path_input",
        )
        c1, c2, c3 = st.columns(3)
        with c1:
            if st.button("▶ Run Phase 1", key="run_p1_dash", use_container_width=True):
                if apk_input and apk_input.strip():
                    _start_phase([sys.executable, "static_analysis.py", apk_input.strip()], "Phase 1")
                    st.rerun()
                else:
                    st.warning("Enter an APK path first.", icon="⚠️")
        with c2:
            if st.button("↺ Reload Sessions", key="reload_ui_dash", use_container_width=True):
                st.cache_data.clear()
                st.session_state.selected_session_name = None
                st.rerun()
        with c3:
            # Download report button
            if selected_session and _DOCX_AVAILABLE:
                docx_bytes = build_word_report(
                    static_report, sentry_report, explorer_report, phase4_report, selected_session
                )
                if docx_bytes:
                    fname = f"threat_report_{selected_session.name}.docx"
                    st.download_button(
                        "📥 Download Report (.docx)",
                        data=docx_bytes,
                        file_name=fname,
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key="dl_report_dash",
                        use_container_width=True,
                    )
            elif not _DOCX_AVAILABLE:
                st.caption("Install `python-docx` to enable report download.")

        # Phase runner output (if Phase 1 is running from this button)
        if _phase_running() or st.session_state.phase_done:
            finished = _poll_phase()
            if finished and not st.session_state.phase_done:
                st.session_state.phase_done = True
                st.cache_data.clear()
                st.session_state.selected_session_name = None
            lines = st.session_state.phase_output[-80:]
            if lines:
                status_color = "var(--red)" if st.session_state.phase_error else "var(--green)" if st.session_state.phase_done else "var(--cyan)"
                status_text  = "FAILED" if st.session_state.phase_error else "COMPLETE" if st.session_state.phase_done else "RUNNING…"
                lines_html = "".join(
                    f'<div style="font-family:var(--font-mono);font-size:11px;line-height:1.6;">{_escape_html(l)}</div>'
                    for l in lines
                )
                st.markdown(f"""
                <div style="background:#060a0e;border:1px solid var(--border);border-radius:4px;
                            padding:10px 14px;margin-top:8px;max-height:200px;overflow-y:auto;">
                  <div style="display:flex;justify-content:space-between;margin-bottom:6px;
                              border-bottom:1px solid var(--border);padding-bottom:4px;">
                    <span style="font-family:var(--font-display);font-size:10px;color:var(--text-dim);">PHASE 1 OUTPUT</span>
                    <span style="font-family:var(--font-mono);font-size:10px;color:{status_color};">{status_text}</span>
                  </div>
                  {lines_html}
                </div>""", unsafe_allow_html=True)
                if _phase_running():
                    st.rerun()

    st.markdown('</div>', unsafe_allow_html=True)

    # ── Automatic next-phase runner ───────────────────────────────────────────
    st.markdown('<div style="padding:0 28px;">', unsafe_allow_html=True)
    if static_report and not sentry_report:
        _p1_path  = str(selected_session / "static_report.json") if selected_session else ""
        _dur      = int(os.environ.get("ANALYSIS_DURATION", "180"))
        _render_phase_runner(
            "Phase 2",
            [sys.executable, "sentry.py", "--report", _p1_path, "--duration", str(_dur)],
        )
    elif sentry_report and not explorer_report:
        _se_path = str(selected_session / "sentry_report.json") if selected_session else ""
        _mx      = int(os.environ.get("EXPLORER_MAX_STATES", "80"))
        _render_phase_runner(
            "Phase 3",
            [sys.executable, "explorer.py", "--sentry", _se_path, "--states", str(_mx)],
        )
    elif explorer_report and not phase4_report:
        _render_phase_runner(
            "Phase 4",
            [sys.executable, "llm_analysis.py", str(selected_session)],
        )
    st.markdown('</div>', unsafe_allow_html=True)

    # ── Metrics ───────────────────────────────────────────────────────────────
    st.markdown('<div style="padding:0 28px;">', unsafe_allow_html=True)
    col1, col2, col3, col4, col5 = st.columns(5)
    for col, lbl, val, color in zip(
        [col1,col2,col3,col4,col5],
        ["PERM RISK SCORE","HARDCODED IOCs","NET CONNECTIONS","EXFIL TYPES","UI STATES"],
        [str(perm_score), str(total_ind), str(net_conns), str(len(exfil_list)), str(states_vis)],
        ["var(--red)" if perm_score>30 else "var(--text-secondary)",
         "var(--orange)" if total_ind>0 else "var(--text-secondary)",
         "var(--yellow)" if net_conns>0 else "var(--text-secondary)",
         "var(--magenta)" if exfil_list else "var(--text-secondary)",
         "var(--cyan)" if states_vis>0 else "var(--text-secondary)"],
    ):
        with col:
            st.markdown(f'<div class="metric-box"><div class="metric-val" style="color:{color};font-size:26px;">{val}</div><div class="metric-lbl">{lbl}</div></div>', unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)
    col_main, col_right = st.columns([2, 1])

    with col_main:
        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("📦", "TARGET APK")
        pkg      = _safe_get(S, "package_name", default="") or _safe_get(S, "package", default="—")
        apk_name = _safe_get(S, "apk_name", default="—")
        apk_size = _safe_get(S, "apk_size_mb", default=0)
        main_act = _safe_get(S, "main_activity", default="—")
        min_sdk  = _safe_get(S, "min_sdk", default=0)
        tgt_sdk  = _safe_get(S, "target_sdk", default=0)
        # PIDs come from sentry (Phase 2 owns launch)
        all_pids = _safe_get(SE, "all_pids", default=[]) or _safe_get(S, "all_pids", default=[])
        pid_disp = ", ".join(str(p) for p in all_pids) if all_pids else "—"
        sdk_color= "var(--yellow)" if tgt_sdk and tgt_sdk < 30 else "var(--text-primary)"
        st.markdown(
            kv("Package",           pkg,      "var(--cyan)") +
            kv("APK File",          apk_name) +
            kv("Size",              f"{apk_size} MB" if apk_size else "—") +
            kv("Main Activity",     main_act, "var(--text-secondary)") +
            kv("Min / Target SDK",  f"{min_sdk} / {tgt_sdk}" if tgt_sdk else "—", sdk_color) +
            kv("PID(s)",            pid_disp, "var(--green)"),
            unsafe_allow_html=True)
        st.markdown('</div>', unsafe_allow_html=True)

        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("⚡", "TOP ABUSED PERMISSIONS — PHASE 2")
        abuse = _safe_get(s_sum, "abuse_scores", default={})
        if abuse and isinstance(abuse, dict):
            try:
                mx = max(abuse.values()) or 1
            except (ValueError, TypeError):
                mx = 1
            for perm, score in sorted(abuse.items(), key=lambda x: -(x[1] if isinstance(x[1], (int,float)) else 0))[:6]:
                try:
                    pct  = float(score) / mx
                    fill = "var(--red)" if score>=80 else "var(--orange)" if score>=60 else "var(--yellow)" if score>=40 else "var(--cyan)"
                    st.markdown(f'<div class="perm-row"><div class="perm-name">{_escape_html(perm)}</div><div class="perm-bar-bg"><div class="perm-bar-fill" style="width:{pct*100:.0f}%;background:{fill};"></div></div><div class="perm-score" style="color:{fill};">{score}</div></div>', unsafe_allow_html=True)
                except Exception:
                    continue
        else:
            no_data("Phase 2 not yet run")
        st.markdown('</div>', unsafe_allow_html=True)

    with col_right:
        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("📤", "EXFILTRATION & DROPPERS")
        exfil   = _safe_get(s_sum, "exfiltration_types", default=[])
        dropper = _safe_get(s_sum, "dropper_chain", default=[])
        packer  = _safe_get(s_sum, "packer_chain", default=[])
        remote  = _safe_get(s_sum, "unique_remote_ips", default=[])
        if exfil or dropper or packer or remote:
            if exfil:
                st.markdown('<div style="font-family:var(--font-display);font-size:9px;letter-spacing:2px;color:var(--text-dim);margin-bottom:6px;">EXFIL TYPES</div>', unsafe_allow_html=True)
                st.markdown(" ".join(f'<span class="ioc-type file">{_escape_html(t).upper()}</span>' for t in exfil), unsafe_allow_html=True)
                st.markdown("<br>", unsafe_allow_html=True)
            if dropper:
                st.markdown('<div style="font-family:var(--font-display);font-size:9px;letter-spacing:2px;color:var(--text-dim);margin:8px 0 4px;">DROPPER CHAIN</div>', unsafe_allow_html=True)
                for d in dropper:
                    _d_str = (d.get("pkg_name") or d.get("package_name") or str(d)) if isinstance(d, dict) else str(d)
                    st.markdown(f'<div style="font-family:var(--font-mono);font-size:11px;color:var(--magenta);">⬇ {_escape_html(_d_str)}</div>', unsafe_allow_html=True)
            if packer:
                st.markdown('<div style="font-family:var(--font-display);font-size:9px;letter-spacing:2px;color:var(--text-dim);margin:8px 0 4px;">PACKER CHAIN</div>', unsafe_allow_html=True)
                for pk in packer:
                    _pk_str = (pk.get("pkg_name") or pk.get("package_name") or str(pk)) if isinstance(pk, dict) else str(pk)
                    st.markdown(f'<div style="font-family:var(--font-mono);font-size:11px;color:var(--orange);">⬡ {_escape_html(_pk_str)}</div>', unsafe_allow_html=True)
            if remote:
                st.markdown('<div style="font-family:var(--font-display);font-size:9px;letter-spacing:2px;color:var(--text-dim);margin:10px 0 4px;">REMOTE IPs</div>', unsafe_allow_html=True)
                for ip in remote:
                    st.markdown(f'<div style="font-family:var(--font-mono);font-size:11px;color:var(--yellow);padding:3px 0;border-bottom:1px solid var(--border);">{_escape_html(str(ip))}</div>', unsafe_allow_html=True)
        else:
            no_data("Phase 2 not yet run")
        st.markdown('</div>', unsafe_allow_html=True)

        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("🎯", "HARDCODED IOCs — PHASE 1")
        inds = _safe_get(summ, "indicators", default={})
        ioc_rows = ""
        for url in _safe_get(inds, "urls", default=[]):
            ioc_rows += f'<tr><td><span class="ioc-type url">URL</span></td><td style="font-size:10px;word-break:break-all;">{_escape_html(str(url))}</td></tr>'
        for ip in _safe_get(inds, "ips", default=[]):
            ioc_rows += f'<tr><td><span class="ioc-type ip">IP</span></td><td style="font-size:10px;">{_escape_html(str(ip))}</td></tr>'
        for dom in _safe_get(inds, "domains", default=[]):
            ioc_rows += f'<tr><td><span class="ioc-type domain">DOM</span></td><td style="font-size:10px;">{_escape_html(str(dom))}</td></tr>'
        if ioc_rows:
            st.markdown(f'<table class="ioc-table"><thead><tr><th>TYPE</th><th>VALUE</th></tr></thead><tbody>{ioc_rows}</tbody></table>', unsafe_allow_html=True)
        else:
            no_data("Phase 1 not yet run")
        st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('</div>', unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════════
#  STATIC ANALYSIS
# ══════════════════════════════════════════════════════════════════════════════

elif "STATIC" in page:
    summ   = _safe_get(S, "summary", default={})
    status = ("COMPLETE ✓","complete") if static_report else ("NOT RUN","pending")

    st.markdown(f"""
    <div class="top-bar">
      <div><span class="brand">Dyn<span>Mal</span>Tool</span>
        <span style="font-family:var(--font-mono);font-size:11px;color:var(--cyan);margin-left:16px;">// STATIC_ANALYSIS · PHASE_1</span>
      </div>
      <span class="phase-badge {status[1]}">{status[0]}</span>
    </div>""", unsafe_allow_html=True)

    if static_report and not sentry_report:
        _p1_path = str(selected_session / "static_report.json") if selected_session else ""
        _dur     = int(os.environ.get("ANALYSIS_DURATION", "180"))
        _render_phase_runner(
            "Phase 2",
            [sys.executable, "sentry.py", "--report", _p1_path, "--duration", str(_dur)],
        )

    st.markdown('<div style="padding:20px 28px;">', unsafe_allow_html=True)
    if not static_report:
        st.warning("No `static_report.json` found — use New Scan on the Dashboard to run Phase 1.", icon="⚠️")

    tgt_sdk = _safe_get(S, "target_sdk", default=0)
    try:
        tgt_sdk = int(tgt_sdk)
    except (TypeError, ValueError):
        tgt_sdk = 0

    col1,col2,col3,col4 = st.columns(4)
    for col, lbl, val, color in zip(
        [col1,col2,col3,col4],
        ["PACKAGE","APK SIZE","TARGET SDK","MIN SDK"],
        [_safe_get(S,"package",default="—"),
         f"{_safe_get(S,'apk_size_mb',default='—')} MB",
         str(tgt_sdk) if tgt_sdk else "—",
         str(_safe_get(S,"min_sdk",default="—"))],
        ["var(--cyan)","var(--text-primary)",
         "var(--yellow)" if tgt_sdk and tgt_sdk<30 else "var(--text-primary)",
         "var(--text-primary)"],
    ):
        with col:
            st.markdown(f'<div class="metric-box"><div class="metric-val" style="color:{color};font-size:14px;word-break:break-all;">{_escape_html(str(val))}</div><div class="metric-lbl">{lbl}</div></div>', unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)
    col_left, col_right = st.columns(2)

    with col_left:
        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("🔑", "PERMISSION RISK TAXONOMY")
        perm_score = _safe_get(summ, "permission_risk_score", default=0)
        top_perms  = _safe_get(summ, "top_risk_permissions", default=[])
        st.markdown(f'<div style="font-family:var(--font-display);font-size:12px;color:var(--text-secondary);margin-bottom:12px;">Total Risk Score: <span style="color:var(--red);font-size:18px;font-weight:700;">{perm_score}</span></div>', unsafe_allow_html=True)
        if top_perms:
            for p in top_perms:
                if not isinstance(p, dict):
                    continue
                try:
                    score = p.get("risk_score", 0)
                    pct   = float(score) / 10
                    fill  = "var(--red)" if score>=9 else "var(--orange)" if score>=7 else "var(--yellow)"
                    name  = _escape_html(str(p.get("name","—")))
                    st.markdown(f'<div class="perm-row"><div class="perm-name">{name}</div><div class="perm-bar-bg"><div class="perm-bar-fill" style="width:{pct*100:.0f}%;background:{fill};"></div></div><div class="perm-score" style="color:{fill};">{score}</div></div>', unsafe_allow_html=True)
                except Exception:
                    continue
        else:
            no_data("No dangerous permissions detected")
        st.markdown('</div>', unsafe_allow_html=True)

        cert = _safe_get(summ, "certificate", default={})
        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("🏅", "CERTIFICATE ANALYSIS")
        cert_html = ""
        if _safe_get(cert, "self_signed", default=False):
            cert_html += f'<div style="display:flex;align-items:center;justify-content:space-between;padding:7px 0;border-bottom:1px solid var(--border);"><span style="font-size:13px;color:var(--text-secondary);">Self-Signed</span>{sev_badge("high")}</div>'
        if _safe_get(cert, "weak_algorithm", default=False):
            alg = _escape_html(str(_safe_get(cert,"algorithm",default="—")))
            cert_html += f'<div style="display:flex;align-items:center;justify-content:space-between;padding:7px 0;border-bottom:1px solid var(--border);"><span style="font-size:13px;color:var(--text-secondary);">Weak Algo: {alg}</span>{sev_badge("high")}</div>'
        if _safe_get(cert, "expired", default=False):
            cert_html += f'<div style="display:flex;align-items:center;justify-content:space-between;padding:7px 0;border-bottom:1px solid var(--border);"><span style="font-size:13px;color:var(--text-secondary);">Certificate Expired</span>{sev_badge("high")}</div>'
        if not _safe_get(cert, "expired", default=False) and _safe_get(cert,"algorithm",default="—") != "—":
            cert_html += f'<div style="display:flex;align-items:center;justify-content:space-between;padding:7px 0;border-bottom:1px solid var(--border);"><span style="font-size:13px;color:var(--text-secondary);">Not Expired</span>{sev_badge("good")}</div>'
        if cert_html:
            st.markdown(cert_html, unsafe_allow_html=True)
            fp = _escape_html(str(_safe_get(cert,"fingerprint",default="—")))
            st.markdown(f'<div style="font-family:var(--font-mono);font-size:10px;color:var(--text-dim);margin-top:10px;">SHA-256: {fp}</div>', unsafe_allow_html=True)
        else:
            no_data("No certificate data")
        st.markdown('</div>', unsafe_allow_html=True)

    with col_right:
        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("📄", "MANIFEST SECURITY ISSUES")
        issues = _safe_get(summ, "manifest_issues", default=[])
        if issues:
            for issue in issues:
                if not isinstance(issue, dict):
                    continue
                try:
                    sev_val = issue.get("severity","info")
                    accent  = {"high":"red","warning":"yel","info":"cyan","good":"green"}.get(sev_val,"cyan")
                    title_i = _escape_html(str(issue.get("title","—")))
                    desc_i  = _escape_html(str(issue.get("description","")))
                    comp_i  = _escape_html(str(issue.get("component","")))
                    st.markdown(f"""
                    <div class="dmt-card dmt-card-accent-{accent}" style="margin-bottom:8px;">
                      <div style="display:flex;align-items:center;justify-content:space-between;margin-bottom:4px;">
                        <span style="font-family:var(--font-display);font-size:12px;font-weight:700;color:var(--text-primary);">{title_i}</span>
                        {sev_badge(sev_val)}
                      </div>
                      <div style="font-size:12px;color:var(--text-secondary);line-height:1.4;">{desc_i}</div>
                      <div style="font-family:var(--font-mono);font-size:10px;color:var(--text-dim);margin-top:4px;">{comp_i}</div>
                    </div>""", unsafe_allow_html=True)
                except Exception:
                    continue
        else:
            no_data("No manifest issues detected")
        st.markdown('</div>', unsafe_allow_html=True)

        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("🔎", "CODE PATTERN SCAN")
        patterns = _safe_get(summ, "code_patterns", default={})
        threats  = _safe_get(patterns, "threats", default=[])
        defences = _safe_get(patterns, "defences", default=[])
        if threats:
            for hit in threats:
                if not isinstance(hit, dict):
                    continue
                try:
                    desc = str(hit.get("description",""))
                    st.markdown(f"""
                    <div style="display:flex;align-items:center;justify-content:space-between;padding:7px 0;border-bottom:1px solid var(--border);">
                      <div>
                        <div style="font-family:var(--font-body);font-size:13px;color:var(--text-secondary);">{_escape_html(str(hit.get("title","—")))}</div>
                        <div style="font-family:var(--font-mono);font-size:10px;color:var(--text-dim);">{_escape_html(desc[:80])}{"…" if len(desc)>80 else ""}</div>
                      </div>
                      {sev_badge(hit.get("severity","info"))}
                    </div>""", unsafe_allow_html=True)
                except Exception:
                    continue
        else:
            no_data("No threat patterns detected")
        if defences:
            st.markdown('<div style="margin-top:10px;font-family:var(--font-display);font-size:10px;letter-spacing:2px;color:var(--green);">DEFENCE INDICATORS</div>', unsafe_allow_html=True)
            for d in defences:
                if isinstance(d, dict):
                    st.markdown(f'<div style="display:flex;align-items:center;justify-content:space-between;padding:6px 0;border-bottom:1px solid var(--border);"><span style="font-size:13px;color:var(--text-secondary);">{_escape_html(str(d.get("title","—")))}</span>{sev_badge("good")}</div>', unsafe_allow_html=True)
        st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
    section_title("📡", "HARDCODED INDICATORS")
    inds  = _safe_get(summ, "indicators", default={})
    col_u, col_i, col_d = st.columns(3)
    for col, lbl, items, cls in [
        (col_u,"URLs",    _safe_get(inds,"urls",default=[]),    "url"),
        (col_i,"IPs",     _safe_get(inds,"ips",default=[]),     "ip"),
        (col_d,"DOMAINS", _safe_get(inds,"domains",default=[]), "domain"),
    ]:
        with col:
            st.markdown(f'<div style="font-family:var(--font-display);font-size:9px;letter-spacing:2px;color:var(--text-dim);text-transform:uppercase;margin-bottom:8px;">{lbl} ({len(items)})</div>', unsafe_allow_html=True)
            if items:
                for item in items:
                    st.markdown(f'<div style="font-family:var(--font-mono);font-size:11px;color:var(--text-secondary);padding:4px 0;border-bottom:1px solid var(--border);word-break:break-all;">{_escape_html(str(item))}</div>', unsafe_allow_html=True)
            else:
                st.markdown('<div style="font-family:var(--font-mono);font-size:11px;color:var(--text-dim);">none</div>', unsafe_allow_html=True)
    st.markdown('</div>', unsafe_allow_html=True)
    st.markdown('</div>', unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════════
#  DYNAMIC SENTRY
# ══════════════════════════════════════════════════════════════════════════════

elif "SENTRY" in page:
    se_sum = _safe_get(SE, "summary", default={})
    status = ("COMPLETE ✓","complete") if sentry_report else ("NOT RUN","pending")

    st.markdown(f"""
    <div class="top-bar">
      <div><span class="brand">Dyn<span>Mal</span>Tool</span>
        <span style="font-family:var(--font-mono);font-size:11px;color:var(--cyan);margin-left:16px;">// DYNAMIC_SENTRY · PHASE_2</span>
      </div>
      <span class="phase-badge {status[1]}">{status[0]}</span>
    </div>""", unsafe_allow_html=True)

    # ── Phase 2 runner — always available on sentry page ─────────────────────
    st.markdown('<div style="padding:12px 28px 0;">', unsafe_allow_html=True)
    with st.expander("▶ Run / Re-Run Phase 2 (Dynamic Sentry)", expanded=not sentry_report):
        _c1, _c2 = st.columns(2)
        with _c1:
            _no_explore = st.checkbox("--no-explore (monitor only)", key="p2_no_explore", value=False)
            _dur2 = st.number_input("Duration (s)", value=int(os.environ.get("ANALYSIS_DURATION","180")),
                                     min_value=30, max_value=3600, label_visibility="visible", key="p2_dur")
        with _c2:
            _direct_apk = st.text_input("--apk <package> (direct mode — bypasses Phase 1 report)",
                                         placeholder="com.evil.package", label_visibility="visible", key="p2_apk")
            st.caption("Enter package name for direct mode. Leave blank to use Phase 1 report.")

        _p1_path = str(selected_session / "static_report.json") if selected_session else ""
        if _direct_apk and _direct_apk.strip():
            _p2_cmd = [sys.executable, "sentry.py", "--apk", _direct_apk.strip(), "--duration", str(_dur2)]
        elif static_report and selected_session:
            _p2_cmd = [sys.executable, "sentry.py", "--report", _p1_path, "--duration", str(_dur2)]
        else:
            _p2_cmd = None

        if _no_explore and _p2_cmd:
            _p2_cmd.append("--no-explore")

        if _p2_cmd:
            _render_phase_runner("Phase 2", _p2_cmd)
        else:
            st.warning("Enter a package name above (direct mode) or run Phase 1 first.", icon="⚠️")

    # Also offer Phase 3 / 4 shortcuts if prior phases done
    if sentry_report and not explorer_report:
        _render_phase_runner("Phase 3",
            [sys.executable, "explorer.py",
             "--sentry", str(selected_session / "sentry_report.json") if selected_session else ""])
    elif explorer_report and not phase4_report:
        _render_phase_runner("Phase 4",
            [sys.executable, "llm_analysis.py", str(selected_session)])
    st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('<div style="padding:20px 28px;">', unsafe_allow_html=True)
    if not sentry_report:
        st.warning("No sentry report found — run Phase 2 first.", icon="⚠️")

    # Key reads — real sentry_report uses top-level fields, not only summary
    logcat_ev = (
        _safe_get(se_sum, "total_log_events", default=0) or
        _safe_get(se_sum, "logcat_events", default=0) or
        len(_safe_get(SE, "log_events", default=[]))
    )
    _net_conn_list = _safe_get(SE, "network_connections", default=[])
    net_conns = (
        _safe_get(se_sum, "total_network_connections", default=0) or
        _safe_get(se_sum, "network_connections", default=0) or
        len(_net_conn_list)
    )
    exfil_cnt = len(_safe_get(se_sum, "exfiltration_types", default=[]))
    # derive duration from timestamps when field absent
    duration  = _safe_get(SE, "duration_seconds", default=0)
    if not duration:
        try:
            from datetime import datetime as _dt
            _ss = _safe_get(SE, "session_start", default="")
            _se = _safe_get(SE, "session_end", default="")
            if _ss and _se:
                duration = round((_dt.fromisoformat(_se) - _dt.fromisoformat(_ss)).total_seconds())
        except Exception:
            pass
    root_det  = bool(_safe_get(se_sum, "root_detected", default=False) or _safe_get(SE, "root_detected", default=False))
    crashes   = _safe_get(se_sum, "crashes", default=0)
    dialogs   = (
        _safe_get(se_sum, "startup_dialogs_handled", default=0) or
        _safe_get(SE, "startup_dialogs_handled", default=0)
    )

    col1, col2, col3, col4, col5 = st.columns(5)
    for col, lbl, val, color in zip(
        [col1,col2,col3,col4,col5],
        ["LOGCAT EVENTS","NETWORK CONNS","EXFIL TYPES","DURATION (s)","STARTUP DIALOGS"],
        [str(logcat_ev),str(net_conns),str(exfil_cnt),str(duration),str(dialogs)],
        ["var(--cyan)" if logcat_ev else "var(--text-dim)",
         "var(--yellow)" if net_conns else "var(--text-dim)",
         "var(--red)" if exfil_cnt else "var(--text-dim)",
         "var(--text-primary)",
         "var(--green)" if dialogs else "var(--text-dim)"],
    ):
        with col:
            st.markdown(f'<div class="metric-box"><div class="metric-val" style="color:{color};">{val}</div><div class="metric-lbl">{lbl}</div></div>', unsafe_allow_html=True)

    if root_det or crashes:
        alert_html = ""
        if root_det:
            alert_html += ('<div style="display:flex;align-items:center;gap:12px;padding:10px 16px;margin:12px 0 8px;'
                           'background:rgba(255,68,68,0.08);border:1px solid rgba(255,68,68,0.4);border-left:3px solid var(--red);border-radius:3px;">'
                           '<span style="font-size:16px;">🚨</span>'
                           '<span style="font-family:var(--font-display);font-size:12px;font-weight:700;letter-spacing:1.5px;color:var(--red);">'
                           'ROOT DETECTION FIRED</span>'
                           '<span style="font-family:var(--font-mono);font-size:11px;color:var(--text-secondary);margin-left:8px;">'
                           '— App may have self-terminated or wiped data</span></div>')
        if crashes:
            alert_html += (f'<div style="display:flex;align-items:center;gap:12px;padding:10px 16px;margin-bottom:8px;'
                           f'background:rgba(255,196,0,0.07);border:1px solid rgba(255,196,0,0.35);border-left:3px solid var(--yellow);border-radius:3px;">'
                           f'<span style="font-size:16px;">⚠️</span>'
                           f'<span style="font-family:var(--font-display);font-size:12px;font-weight:700;letter-spacing:1.5px;color:var(--yellow);">'
                           f'{crashes} CRASH{"ES" if crashes != 1 else ""} DETECTED</span>'
                           f'<span style="font-family:var(--font-mono);font-size:11px;color:var(--text-secondary);margin-left:8px;">'
                           f'— Recovered by CrashRecovery thread</span></div>')
        st.markdown(alert_html, unsafe_allow_html=True)

    # ── Packer chain alert ────────────────────────────────────────────────────
    packer_chain = _safe_get(se_sum, "packer_chain", default=[])
    if packer_chain:
        pc_items = []
        for p in packer_chain:
            if isinstance(p, dict):
                pkg  = p.get("pkg_name") or p.get("package_name") or "unknown"
                sha  = str(p.get("sha256",""))[:12]
                meth = p.get("pull_method","")
                idx  = p.get("index","")
                pc_items.append(f"#{idx} {pkg}" + (f" [{meth}]" if meth else "") + (f" {sha}…" if sha else ""))
            else:
                pc_items.append(str(p))
        pc_str = " → ".join(_escape_html(s) for s in pc_items)
        st.markdown(
            f'<div style="display:flex;align-items:center;gap:12px;padding:10px 16px;margin-bottom:8px;'
            f'background:rgba(255,140,0,0.08);border:1px solid rgba(255,140,0,0.35);border-left:3px solid var(--orange);border-radius:3px;">'
            f'<span style="font-size:16px;">📦</span>'
            f'<div><span style="font-family:var(--font-display);font-size:12px;font-weight:700;color:var(--orange);">PACKER CHAIN DETECTED</span>'
            f'<div style="font-family:var(--font-mono);font-size:11px;color:var(--text-secondary);margin-top:3px;">{pc_str}</div></div>'
            f'</div>',
            unsafe_allow_html=True,
        )

    st.markdown("<br>", unsafe_allow_html=True)
    col_left, col_right = st.columns(2)

    with col_left:
        threads = _safe_get(SE, "threads", default=[])
        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("🔀", f"MONITOR THREADS ({len(threads)})")
        if threads:
            for t in threads:
                if not isinstance(t, dict):
                    continue
                dcls = {"active":"active","idle":"idle","error":"error"}.get(t.get("status","idle"),"idle")
                tname = _escape_html(str(t.get("name","—")))
                tevt  = t.get("events", 0)
                st.markdown(f'<div class="thread-row"><div class="thread-dot {dcls}"></div><div class="thread-name">{tname}</div><div class="thread-events">{tevt} events</div></div>', unsafe_allow_html=True)
        else:
            no_data("No thread data")
        st.markdown('</div>', unsafe_allow_html=True)

        net_list = _safe_get(SE, "network_connections", default=[])
        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("🌐", "NETWORK CONNECTIONS")
        if net_list:
            rows = ""
            for conn in net_list:
                if not isinstance(conn, dict):
                    continue
                flag  = str(conn.get("flag","—"))
                fc    = {"C2":"var(--red)","EXFIL":"var(--magenta)","BEACON":"var(--orange)","DNS":"var(--text-dim)"}.get(flag,"var(--cyan)")
                # real sentry_report uses remote_ip / remote_port / protocol
                ip_val    = conn.get("remote_ip") or conn.get("ip") or "—"
                port_val  = conn.get("remote_port") or conn.get("port") or "—"
                proto_val = conn.get("protocol") or conn.get("proto") or "—"
                domain_val = conn.get("domain", "")
                # show domain if available, otherwise just IP
                ip_display = f"{ip_val}" + (f" ({domain_val})" if domain_val else "")
                rows += (f'<tr>'
                         f'<td style="font-family:var(--font-mono);font-size:11px;">{_escape_html(ip_display)}</td>'
                         f'<td style="font-family:var(--font-mono);font-size:11px;">{_escape_html(str(port_val))}</td>'
                         f'<td style="font-family:var(--font-mono);font-size:11px;">{_escape_html(str(proto_val))}</td>'
                         f'<td><span style="font-family:var(--font-display);font-size:10px;font-weight:700;color:{fc};">{_escape_html(flag)}</span></td>'
                         f'</tr>')
            st.markdown(f'<table class="ioc-table"><thead><tr><th>IP</th><th>PORT</th><th>PROTO</th><th>FLAG</th></tr></thead><tbody>{rows}</tbody></table>', unsafe_allow_html=True)
        else:
            no_data("No new external connections detected")
        st.markdown('</div>', unsafe_allow_html=True)

    with col_right:
        abuse = _safe_get(se_sum, "abuse_scores", default={})
        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("⚡", "PERMISSION ABUSE SCORES")
        if abuse and isinstance(abuse, dict):
            try:
                mx = max(abuse.values()) or 1
            except (ValueError, TypeError):
                mx = 1
            for perm, score in sorted(abuse.items(), key=lambda x: -(x[1] if isinstance(x[1],(int,float)) else 0)):
                try:
                    pct  = float(score) / mx
                    fill = "var(--red)" if score>=80 else "var(--orange)" if score>=60 else "var(--yellow)" if score>=40 else "var(--cyan)"
                    st.markdown(f'<div class="perm-row"><div class="perm-name">{_escape_html(perm)}</div><div class="perm-bar-bg"><div class="perm-bar-fill" style="width:{pct*100:.0f}%;background:{fill};"></div></div><div class="perm-score" style="color:{fill};">{score}</div></div>', unsafe_allow_html=True)
                except Exception:
                    continue
        else:
            no_data("No permission events recorded")
        st.markdown('</div>', unsafe_allow_html=True)

        exfil   = _safe_get(se_sum, "exfiltration_types", default=[])
        dropper = _safe_get(se_sum, "dropper_chain", default=[])
        remote  = _safe_get(se_sum, "unique_remote_ips", default=[])
        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("📤", "EXFILTRATION & DROPPER CHAIN")
        if exfil or dropper or packer_chain or remote:
            if exfil:
                st.markdown('<div style="font-family:var(--font-display);font-size:9px;letter-spacing:2px;color:var(--text-dim);margin-bottom:6px;">EXFIL TYPES</div>', unsafe_allow_html=True)
                st.markdown(" ".join(f'<span class="ioc-type file">{_escape_html(t).upper()}</span>' for t in exfil), unsafe_allow_html=True)
                st.markdown("<br><br>", unsafe_allow_html=True)
            if dropper:
                st.markdown('<div style="font-family:var(--font-display);font-size:9px;letter-spacing:2px;color:var(--text-dim);margin-bottom:6px;">DROPPER CHAIN</div>', unsafe_allow_html=True)
                for d in dropper:
                    _d_str = (d.get("pkg_name") or d.get("package_name") or str(d)) if isinstance(d, dict) else str(d)
                    st.markdown(f'<div style="font-family:var(--font-mono);font-size:12px;color:var(--magenta);">⬇ {_escape_html(_d_str)}</div>', unsafe_allow_html=True)
                st.markdown("<br>", unsafe_allow_html=True)
            if packer_chain:
                st.markdown('<div style="font-family:var(--font-display);font-size:9px;letter-spacing:2px;color:var(--text-dim);margin-bottom:6px;">PACKER CHAIN</div>', unsafe_allow_html=True)
                for pk in packer_chain:
                    _pk_str = (pk.get("pkg_name") or pk.get("package_name") or str(pk)) if isinstance(pk, dict) else str(pk)
                    st.markdown(f'<div style="font-family:var(--font-mono);font-size:12px;color:var(--orange);">⬡ {_escape_html(_pk_str)}</div>', unsafe_allow_html=True)
                st.markdown("<br>", unsafe_allow_html=True)
            if remote:
                st.markdown('<div style="font-family:var(--font-display);font-size:9px;letter-spacing:2px;color:var(--text-dim);margin-bottom:6px;">REMOTE IPs</div>', unsafe_allow_html=True)
                for ip in remote:
                    st.markdown(f'<div style="font-family:var(--font-mono);font-size:12px;color:var(--yellow);padding:3px 0;border-bottom:1px solid var(--border);">{_escape_html(str(ip))}</div>', unsafe_allow_html=True)
        else:
            no_data("No exfiltration or dropper activity detected")
        st.markdown('</div>', unsafe_allow_html=True)

    # ── Sentry errors ─────────────────────────────────────────────────────────
    sentry_errs = _safe_get(SE, "errors", default=[])
    if sentry_errs:
        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("⚠", f"SENTRY ERRORS ({len(sentry_errs)})")
        for err in sentry_errs:
            st.markdown(f'<div style="font-family:var(--font-mono);font-size:11px;color:var(--red);padding:4px 0;border-bottom:1px solid var(--border);">⚠ {_escape_html(str(err))}</div>', unsafe_allow_html=True)
        st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('</div>', unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════════
#  EXPLORER
# ══════════════════════════════════════════════════════════════════════════════

elif "EXPLORER" in page:
    ex_sum = _safe_get(EX, "summary", default={})
    status = ("COMPLETE ✓","complete") if explorer_report else ("NOT RUN","pending")

    st.markdown(f"""
    <div class="top-bar">
      <div><span class="brand">Dyn<span>Mal</span>Tool</span>
        <span style="font-family:var(--font-mono);font-size:11px;color:var(--cyan);margin-left:16px;">// DFS_EXPLORER · PHASE_3</span>
      </div>
      <span class="phase-badge {status[1]}">{status[0]}</span>
    </div>""", unsafe_allow_html=True)

    # ── Phase 3 runner — always available on explorer page ───────────────────
    st.markdown('<div style="padding:12px 28px 0;">', unsafe_allow_html=True)
    with st.expander("▶ Run / Re-Run Phase 3 (DFS Explorer)", expanded=not explorer_report):
        _c1, _c2 = st.columns(2)
        with _c1:
            _p3_apk  = st.text_input("--apk <package> (direct mode — bypasses sentry report)",
                                      placeholder="com.evil.package", label_visibility="visible", key="p3_apk")
            _mx3     = st.number_input("Max states", value=int(os.environ.get("EXPLORER_MAX_STATES","80")),
                                        min_value=10, max_value=500, label_visibility="visible", key="p3_states")
        with _c2:
            _p3_wait = st.checkbox("--wait (wait for manual app launch)", key="p3_wait", value=False)
            _p3_wait_to = st.number_input("Wait timeout (s)", value=120, min_value=30, max_value=600,
                                           label_visibility="visible", key="p3_wait_timeout")
            st.caption("Leave --apk blank to use sentry report from current session.")

        _se_path = str(selected_session / "sentry_report.json") if selected_session else ""
        if _p3_apk and _p3_apk.strip():
            _p3_cmd = [sys.executable, "explorer.py", "--apk", _p3_apk.strip(), "--states", str(_mx3)]
        elif sentry_report and selected_session:
            _p3_cmd = [sys.executable, "explorer.py", "--sentry", _se_path, "--states", str(_mx3)]
        else:
            _p3_cmd = None

        if _p3_wait and _p3_cmd:
            _p3_cmd += ["--wait", "--wait-timeout", str(_p3_wait_to)]

        if _p3_cmd:
            _render_phase_runner("Phase 3", _p3_cmd)
        else:
            st.warning("Enter a package name above (direct mode) or run Phase 2 first.", icon="⚠️")

    if explorer_report and not phase4_report:
        _render_phase_runner("Phase 4",
            [sys.executable, "llm_analysis.py", str(selected_session)])
    st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('<div style="padding:20px 28px;">', unsafe_allow_html=True)
    if not explorer_report:
        st.warning("No explorer report found — run Phase 3 first.", icon="⚠️")

    col1,col2,col3,col4 = st.columns(4)
    for col, lbl, val, color in zip(
        [col1,col2,col3,col4],
        ["STATES VISITED","TOTAL ACTIONS","DEEPEST DEPTH","DIALOGS ACCEPTED"],
        [str(_safe_get(ex_sum,"states_visited",default=0) or _safe_get(EX,"states_visited",default=0)),
         str(_safe_get(ex_sum,"total_actions",default=0) or _safe_get(EX,"total_actions",default=0)),
         str(_safe_get(ex_sum,"deepest_state",default=0) or _safe_get(ex_sum,"deepest_depth",default=0)),
         str(_safe_get(ex_sum,"permission_dialogs_accepted",default=0) or _safe_get(EX,"permission_dialogs_accepted",default=0))],
        ["var(--cyan)","var(--yellow)","var(--magenta)","var(--red)"],
    ):
        with col:
            st.markdown(f'<div class="metric-box"><div class="metric-val" style="color:{color};">{val}</div><div class="metric-lbl">{lbl}</div></div>', unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)
    col_left, col_right = st.columns(2)
    activities = _safe_get(ex_sum, "unique_activities", default=[])

    with col_left:
        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("🗺️", f"UNIQUE ACTIVITIES REACHED ({len(activities)})")
        if activities:
            depth_colors = ["var(--cyan)","var(--green)","var(--yellow)","var(--orange)","var(--magenta)"]
            for i, act in enumerate(activities):
                try:
                    short = str(act).split(".")[-1]
                    color = depth_colors[i % len(depth_colors)]
                    st.markdown(f'<div style="display:flex;align-items:center;gap:10px;padding:8px 0;border-bottom:1px solid var(--border);"><span style="font-family:var(--font-mono);font-size:10px;color:var(--text-dim);">D{i+1}</span><div><div style="font-family:var(--font-display);font-size:13px;font-weight:600;color:{color};">{_escape_html(short)}</div><div style="font-family:var(--font-mono);font-size:9px;color:var(--text-dim);">{_escape_html(str(act))}</div></div></div>', unsafe_allow_html=True)
                except Exception:
                    continue
        else:
            no_data("No activities recorded")
        st.markdown('</div>', unsafe_allow_html=True)

        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("📋", "SESSION DETAILS")
        sess_dir_str  = _safe_get(EX, "session_dir", default="—")
        sess_dir_name = str(sess_dir_str).split("/")[-1] if "/" in str(sess_dir_str) else str(sess_dir_str)
        scr_tier      = _safe_get(ex_sum, "screenshot_tier", default=3) or _safe_get(EX, "screenshot_tier", default=3)
        cr_count      = _safe_get(ex_sum, "crash_recoveries", default=0) or _safe_get(ex_sum, "crashes_recovered", default=0) or _safe_get(EX, "crashes_recovered", default=0)
        forms_filled  = _safe_get(ex_sum, "forms_filled", default=0) or _safe_get(EX, "forms_filled", default=0)
        submits_ok    = _safe_get(ex_sum, "form_submits_succeeded", default=0) or _safe_get(EX, "form_submits_succeeded", default=0)
        submits_tried = _safe_get(ex_sum, "form_submits_attempted", default=0) or _safe_get(EX, "form_submits_attempted", default=0)
        st.markdown(
            kv("Forms Filled",       str(forms_filled), "var(--cyan)") +
            kv("Form Submits",       f"{submits_ok} / {submits_tried} succeeded", "var(--yellow)" if submits_tried else "var(--text-dim)") +
            kv("Screenshot Tier",    f"Tier {scr_tier}", "var(--green)" if scr_tier<3 else "var(--text-dim)") +
            kv("Crash Recoveries",   str(cr_count), "var(--red)" if cr_count else "var(--green)") +
            kv("Session Dir",        _escape_html(sess_dir_name), "var(--text-dim)"),
            unsafe_allow_html=True)
        st.markdown('</div>', unsafe_allow_html=True)

        # Form submit failures
        fail_list = _safe_get(ex_sum, "form_submit_failures", default=[]) or _safe_get(EX, "form_submit_failures", default=[])
        if fail_list:
            st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
            section_title("⚠", f"FORM SUBMIT FAILURES ({len(fail_list)})")
            for fl in fail_list:
                if not isinstance(fl, dict):
                    continue
                try:
                    btn   = _escape_html(str(fl.get("button","—")))
                    ts_f  = _escape_html(str(fl.get("timestamp","—")))
                    retry = fl.get("retries", 0)
                    err_nodes = fl.get("error_nodes", [])
                    err_html  = ""
                    if err_nodes:
                        for en in err_nodes:
                            if isinstance(en, dict):
                                err_html += f'<div style="font-family:var(--font-mono);font-size:10px;color:var(--red);margin-top:4px;padding-left:12px;">→ {_escape_html(str(en.get("text","")))}</div>'
                    st.markdown(
                        f'<div class="dmt-card dmt-card-accent-yel" style="margin-bottom:6px;">'
                        f'<div style="display:flex;align-items:center;justify-content:space-between;margin-bottom:2px;">'
                        f'<span style="font-family:var(--font-display);font-size:12px;font-weight:600;color:var(--text-primary);">{btn}</span>'
                        f'<span style="font-family:var(--font-mono);font-size:10px;color:var(--text-dim);">{ts_f} · {retry} retries</span>'
                        f'</div>{err_html}</div>',
                        unsafe_allow_html=True,
                    )
                except Exception:
                    continue
            st.markdown('</div>', unsafe_allow_html=True)

    with col_right:
        perms_req = _safe_get(ex_sum, "permissions_requested", default=[])
        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("🔐", f"RUNTIME PERMISSION DIALOGS ({len(perms_req)})")
        if perms_req:
            for perm in perms_req:
                short = str(perm).replace("android.permission.","")
                st.markdown(f'<div style="display:flex;align-items:center;justify-content:space-between;padding:8px 0;border-bottom:1px solid var(--border);"><span style="font-family:var(--font-mono);font-size:12px;color:var(--yellow);">{_escape_html(short)}</span><span class="sev-badge sev-high">ACCEPTED</span></div>', unsafe_allow_html=True)
        else:
            no_data("No permission dialogs triggered")
        st.markdown('</div>', unsafe_allow_html=True)

        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("🌲", "DFS ACTIVITY CHAIN")
        if activities:
            tree_html = '<div style="font-family:var(--font-mono);font-size:11px;line-height:2;color:var(--text-secondary);">'
            for i, act in enumerate(activities):
                try:
                    short     = _escape_html(str(act).split(".")[-1])
                    indent    = "&nbsp;" * (i * 4)
                    connector = "" if i == 0 else "└─ " if i == len(activities)-1 else "├─ "
                    cls       = "last" if i == len(activities)-1 else "visited"
                    tree_html += f'{indent}{connector}<span class="state-node {cls}">{short}</span><br>'
                except Exception:
                    continue
            tree_html += '</div>'
            st.markdown(tree_html, unsafe_allow_html=True)
        else:
            no_data("No traversal data")
        st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('</div>', unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════════
#  LLM ANALYSIS  (Phase 4)
# ══════════════════════════════════════════════════════════════════════════════

elif "LLM" in page:
    status     = ("COMPLETE ✓","complete") if phase4_report else ("NOT RUN","pending")
    llm_status = _safe_get(P4, "llm_status", default="not_run")

    # phase4_report nests all verdict data under "llm_verdict"
    _verdict = _safe_get(P4, "llm_verdict", default={}) or {}

    threat_score = _safe_get(_verdict, "threat_score", default=0) or _safe_get(P4, "threat_score", default=0)
    try:
        threat_score = int(threat_score)
    except (TypeError, ValueError):
        threat_score = 0

    score_color = (
        "var(--red)"           if threat_score >= 75 else
        "var(--orange)"        if threat_score >= 50 else
        "var(--yellow)"        if threat_score >= 25 else
        "var(--text-secondary)"
    )

    st.markdown(f"""
    <div class="top-bar">
      <div><span class="brand">Dyn<span>Mal</span>Tool</span>
        <span style="font-family:var(--font-mono);font-size:11px;color:var(--cyan);margin-left:16px;">// LLM_ANALYSIS · PHASE_4</span>
      </div>
      <div style="display:flex;align-items:center;gap:12px;">
        {'<span style="font-family:var(--font-mono);font-size:11px;color:var(--yellow);">⚠ OLLAMA UNAVAILABLE</span>' if llm_status == "ollama_unavailable" else ""}
        <span class="phase-badge {status[1]}">{status[0]}</span>
      </div>
    </div>""", unsafe_allow_html=True)

    if explorer_report and not phase4_report:
        _render_phase_runner("Phase 4",
            [sys.executable, "llm_analysis.py", str(selected_session)])

    st.markdown('<div style="padding:12px 28px 0;">', unsafe_allow_html=True)
    with st.expander("▶ Run / Re-Run Phase 4 (LLM Analysis)", expanded=not phase4_report):
        if selected_session:
            _render_phase_runner("Phase 4 (rerun)",
                [sys.executable, "llm_analysis.py", str(selected_session)])
        else:
            st.warning("Select a session first.", icon="⚠️")
    st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('<div style="padding:20px 28px;">', unsafe_allow_html=True)
    if not phase4_report:
        st.warning("No Phase 4 report found — run Phase 4 (LLM Analysis) first.", icon="⚠️")

    threat_class   = _safe_get(_verdict, "threat_class", default="—") or _safe_get(P4, "threat_class", default="—")
    # confidence is a string ("LOW"/"MEDIUM"/"HIGH") in real report — convert to int %
    _conf_raw      = _safe_get(_verdict, "confidence", default="") or _safe_get(P4, "confidence", default=0)
    if isinstance(_conf_raw, str):
        confidence = {"LOW": 25, "MEDIUM": 55, "HIGH": 85, "CRITICAL": 95}.get(_conf_raw.upper(), 0)
    else:
        try:
            confidence = int(_conf_raw)
        except (TypeError, ValueError):
            confidence = 0
    kill_chain     = _safe_get(_verdict, "kill_chain_stage", default="—") or _safe_get(P4, "kill_chain_stage", default="—")
    # iocs in llm_verdict is a dict of lists: {domains:[], ips:[], urls:[], ...}
    # flatten to a list of {type, value} for the table renderer
    _iocs_dict = _safe_get(_verdict, "iocs", default={}) or {}
    iocs = []
    if isinstance(_iocs_dict, dict):
        for ioc_type, ioc_list in _iocs_dict.items():
            if isinstance(ioc_list, list):
                for v in ioc_list:
                    iocs.append({"type": ioc_type, "value": str(v), "flag": ""})
    elif isinstance(_iocs_dict, list):
        iocs = _iocs_dict
    # also surface network_iocs from phase4_report top level
    for ni in _safe_get(P4, "network_iocs", default=[]):
        if isinstance(ni, dict):
            iocs.append({"type": ni.get("ioc_type", "ip"), "value": ni.get("value", ""), "flag": ni.get("classification", "")})
    data_collected = _safe_get(_verdict, "data_collected", default=[]) or _safe_get(P4, "data_collected", default=[])
    recommended    = _safe_get(_verdict, "recommended_actions", default=[]) or _safe_get(P4, "recommended_actions", default=[])
    evidence       = _safe_get(_verdict, "evidence", default=[]) or _safe_get(P4, "evidence", default=[])
    analyst_notes  = _safe_get(_verdict, "analyst_notes", default="—") or _safe_get(P4, "analyst_notes", default="—")

    col1, col2, col3, col4 = st.columns(4)
    for col, lbl, val, color in zip(
        [col1, col2, col3, col4],
        ["THREAT SCORE", "CONFIDENCE %", "IOC COUNT", "KILL CHAIN STAGE"],
        [str(threat_score), str(confidence), str(len(iocs)), str(kill_chain)],
        [score_color,
         "var(--green)" if confidence >= 70 else "var(--yellow)" if confidence >= 40 else "var(--text-dim)",
         "var(--red)" if iocs else "var(--text-dim)",
         "var(--magenta)" if kill_chain != "—" else "var(--text-dim)"],
    ):
        with col:
            vdisp = _escape_html(str(val))
            st.markdown(
                f'<div class="metric-box">'
                f'<div class="metric-val" style="color:{color};font-size:{"22px" if len(str(val))>6 else "32px"};">{vdisp}</div>'
                f'<div class="metric-lbl">{lbl}</div>'
                f'</div>',
                unsafe_allow_html=True,
            )

    st.markdown("<br>", unsafe_allow_html=True)

    if threat_class and threat_class != "—":
        st.markdown(
            f'<div style="display:flex;align-items:center;gap:16px;padding:14px 20px;margin-bottom:16px;'
            f'background:rgba(255,68,68,0.06);border:1px solid {score_color};border-left:4px solid {score_color};border-radius:3px;">'
            f'<span style="font-size:22px;">🎯</span>'
            f'<div>'
            f'<div style="font-family:var(--font-display);font-size:10px;letter-spacing:2px;color:var(--text-dim);margin-bottom:3px;">THREAT CLASSIFICATION</div>'
            f'<div style="font-family:var(--font-display);font-size:18px;font-weight:700;color:{score_color};">{_escape_html(str(threat_class))}</div>'
            f'</div></div>',
            unsafe_allow_html=True,
        )

    col_left, col_right = st.columns(2)

    with col_left:
        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("🔍", f"EVIDENCE ({len(evidence)})")
        if evidence:
            for ev in evidence:
                try:
                    sev  = ev.get("severity", "info") if isinstance(ev, dict) else "info"
                    desc = ev.get("description", str(ev)) if isinstance(ev, dict) else str(ev)
                    accent = {"high": "red", "warning": "yel", "info": "cyan", "good": "green"}.get(sev, "cyan")
                    desc_safe = _escape_html(str(desc))
                    st.markdown(
                        f'<div class="dmt-card dmt-card-accent-{accent}" style="margin-bottom:6px;">'
                        f'<div style="display:flex;align-items:center;justify-content:space-between;margin-bottom:2px;">'
                        f'<span style="font-family:var(--font-body);font-size:12px;color:var(--text-secondary);">{desc_safe[:120]}{"…" if len(str(desc))>120 else ""}</span>'
                        f'{sev_badge(sev)}'
                        f'</div></div>',
                        unsafe_allow_html=True,
                    )
                except Exception:
                    continue
        else:
            no_data("No evidence extracted")
        st.markdown('</div>', unsafe_allow_html=True)

        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("📝", "ANALYST NOTES")
        if analyst_notes and str(analyst_notes) != "—":
            st.markdown(
                f'<div style="font-family:var(--font-body);font-size:13px;color:var(--text-secondary);'
                f'line-height:1.7;white-space:pre-wrap;">{_escape_html(str(analyst_notes))}</div>',
                unsafe_allow_html=True,
            )
        else:
            no_data("No analyst notes")
        st.markdown('</div>', unsafe_allow_html=True)

    with col_right:
        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("📡", f"IOCs ({len(iocs)})")
        if iocs:
            ioc_rows = ""
            for ioc in iocs:
                try:
                    if isinstance(ioc, dict):
                        ioc_type = str(ioc.get("type", "ioc")).lower()
                        ioc_val  = _escape_html(str(ioc.get("value", "—")))
                        ioc_flag = _escape_html(str(ioc.get("flag", "")))
                    else:
                        ioc_type = "ioc"
                        ioc_val  = _escape_html(str(ioc))
                        ioc_flag = ""
                    type_cls  = ioc_type if ioc_type in ("domain","ip","url","file") else "domain"
                    flag_html = f'<span style="font-family:var(--font-mono);font-size:10px;color:var(--yellow);">{ioc_flag}</span>' if ioc_flag else ""
                    ioc_rows += (f'<tr>'
                                 f'<td><span class="ioc-type {type_cls}">{ioc_type.upper()}</span></td>'
                                 f'<td style="font-size:11px;word-break:break-all;">{ioc_val}</td>'
                                 f'<td>{flag_html}</td>'
                                 f'</tr>')
                except Exception:
                    continue
            st.markdown(
                f'<table class="ioc-table">'
                f'<thead><tr><th>TYPE</th><th>VALUE</th><th>FLAG</th></tr></thead>'
                f'<tbody>{ioc_rows}</tbody>'
                f'</table>',
                unsafe_allow_html=True,
            )
        else:
            no_data("No IOCs extracted")
        st.markdown('</div>', unsafe_allow_html=True)

        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("⚡", f"RECOMMENDED ACTIONS ({len(recommended)})")
        if recommended:
            for i, action in enumerate(recommended, 1):
                action_str = _escape_html(str(action))
                st.markdown(
                    f'<div style="display:flex;align-items:flex-start;gap:12px;padding:8px 0;border-bottom:1px solid var(--border);">'
                    f'<span style="font-family:var(--font-mono);font-size:11px;color:var(--magenta);min-width:20px;">{i:02d}</span>'
                    f'<span style="font-family:var(--font-body);font-size:12px;color:var(--text-secondary);line-height:1.5;">{action_str}</span>'
                    f'</div>',
                    unsafe_allow_html=True,
                )
        else:
            no_data("No recommendations")
        st.markdown('</div>', unsafe_allow_html=True)

        if data_collected:
            st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
            section_title("📦", f"DATA COLLECTED ({len(data_collected)})")
            for item in data_collected:
                st.markdown(
                    f'<div style="font-family:var(--font-mono);font-size:11px;color:var(--yellow);'
                    f'padding:5px 0;border-bottom:1px solid var(--border);">⬇ {_escape_html(str(item))}</div>',
                    unsafe_allow_html=True,
                )
            st.markdown('</div>', unsafe_allow_html=True)

    # ── Suggested follow-up commands ──────────────────────────────────────────
    suggested_cmds = _safe_get(SE, "suggested_commands", default=[])
    if suggested_cmds:
        st.markdown("<br>", unsafe_allow_html=True)
        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("🔗", f"SUGGESTED FOLLOW-UP COMMANDS ({len(suggested_cmds)})")
        cmds_html = '<div style="background:#060a0e;border:1px solid var(--border);border-radius:3px;padding:12px 16px;font-family:var(--font-mono);font-size:11px;line-height:2;">'
        for cmd in suggested_cmds:
            safe_cmd = _escape_html(str(cmd))
            comment  = str(cmd).strip().startswith("#")
            color    = "var(--text-dim)" if comment else "var(--green)"
            cmds_html += f'<div style="color:{color};">{safe_cmd}</div>'
        cmds_html += '</div>'
        st.markdown(cmds_html, unsafe_allow_html=True)
        st.markdown('</div>', unsafe_allow_html=True)

    if llm_status == "ollama_unavailable":
        st.warning(
            "Ollama was unreachable during Phase 4. IOC extraction ran successfully but LLM verdict was skipped. "
            "Start Ollama (`ollama serve`) and re-run Phase 4 to generate a full threat assessment.",
            icon="⚠️",
        )

    st.markdown('</div>', unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════════
#  THREAT REPORT  (Phase 5)
# ══════════════════════════════════════════════════════════════════════════════

elif "THREAT REPORT" in page or "REPORT" in page:
    status = ("COMPLETE ✓","complete") if phase5_report else ("NOT RUN","pending")

    st.markdown(f"""
    <div class="top-bar">
      <div><span class="brand">Dyn<span>Mal</span>Tool</span>
        <span style="font-family:var(--font-mono);font-size:11px;color:var(--cyan);margin-left:16px;">// THREAT_REPORT · PHASE_5</span>
      </div>
      <span class="phase-badge {status[1]}">{status[0]}</span>
    </div>""", unsafe_allow_html=True)

    # Phase 5 runner (when phase5 script exists)
    if phase4_report and not phase5_report:
        if Path("phase5_report.py").exists() or Path("threat_report.py").exists():
            _script = "phase5_report.py" if Path("phase5_report.py").exists() else "threat_report.py"
            _render_phase_runner("Phase 5",
                [sys.executable, _script, str(selected_session)])
        else:
            st.info("Phase 5 (threat_report.py) is not yet implemented. "
                    "Use the Download Report button below to get a Word document of all phases.", icon="ℹ️")

    st.markdown('<div style="padding:20px 28px;">', unsafe_allow_html=True)

    # ── Download Word Report ──────────────────────────────────────────────────
    st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
    section_title("📥", "DOWNLOAD THREAT REPORT")

    if not _DOCX_AVAILABLE:
        st.warning(
            "python-docx is not installed. Run `pip install python-docx` to enable Word report download.",
            icon="⚠️",
        )
    elif not selected_session:
        st.info("Select a session first.", icon="ℹ️")
    else:
        # Generate on the fly — cached by session name in session state
        _cache_key = f"docx_bytes_{selected_session.name}"
        if _cache_key not in st.session_state:
            with st.spinner("Generating report…"):
                try:
                    st.session_state[_cache_key] = build_word_report(
                        static_report, sentry_report, explorer_report, phase4_report, selected_session
                    )
                except Exception as e:
                    st.session_state[_cache_key] = None
                    st.error(f"Report generation failed: {e}", icon="🚨")

        _docx_bytes = st.session_state.get(_cache_key)
        if _docx_bytes:
            fname = f"threat_report_{selected_session.name}.docx"
            st.download_button(
                label="📥 Download Threat Report (.docx)",
                data=_docx_bytes,
                file_name=fname,
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key="dl_report_p5",
            )
            st.markdown(f'<div style="font-family:var(--font-mono);font-size:10px;color:var(--text-dim);margin-top:8px;">Contains: Static Analysis · Dynamic Sentry · UI Explorer · LLM Analysis · IOCs · Recommendations</div>', unsafe_allow_html=True)
        elif _docx_bytes is None and _DOCX_AVAILABLE:
            no_data("No report data available — run at least Phase 1 first")
    st.markdown('</div>', unsafe_allow_html=True)

    # ── Phase 5 report content (if it exists) ────────────────────────────────
    if phase5_report:
        exec_sum = _safe_get(P5, "executive_summary", default="—")
        threat_cls5 = _safe_get(P5, "threat_classification", default="—")
        kill_chain5 = _safe_get(P5, "kill_chain_analysis", default="—")
        recs5 = _safe_get(P5, "recommendations", default=[])
        net_iocs5 = _safe_get(P5, "network_iocs", default=[])
        ui_behav5 = _safe_get(P5, "ui_behaviour", default="—")
        gen_at5 = _safe_get(P5, "generated_at", default="—")

        st.markdown(f'<div style="font-family:var(--font-mono);font-size:9px;color:var(--text-dim);margin-bottom:16px;">Generated: {_escape_html(str(gen_at5))}</div>', unsafe_allow_html=True)

        col_l, col_r = st.columns(2)
        with col_l:
            st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
            section_title("📋", "EXECUTIVE SUMMARY")
            if exec_sum and exec_sum != "—":
                st.markdown(f'<div style="font-family:var(--font-body);font-size:13px;color:var(--text-secondary);line-height:1.7;white-space:pre-wrap;">{_escape_html(str(exec_sum))}</div>', unsafe_allow_html=True)
            else:
                no_data("No summary available — LLM may have been unavailable")
            st.markdown('</div>', unsafe_allow_html=True)

            st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
            section_title("🎯", "THREAT CLASSIFICATION")
            if threat_cls5 and threat_cls5 != "—":
                st.markdown(f'<div style="font-family:var(--font-display);font-size:20px;font-weight:700;color:var(--red);">{_escape_html(str(threat_cls5))}</div>', unsafe_allow_html=True)
            else:
                no_data("Pending LLM analysis")
            st.markdown('</div>', unsafe_allow_html=True)

            st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
            section_title("🔗", "KILL CHAIN ANALYSIS")
            if kill_chain5 and kill_chain5 != "—":
                st.markdown(f'<div style="font-family:var(--font-body);font-size:13px;color:var(--text-secondary);line-height:1.7;white-space:pre-wrap;">{_escape_html(str(kill_chain5))}</div>', unsafe_allow_html=True)
            else:
                no_data("Pending LLM analysis")
            st.markdown('</div>', unsafe_allow_html=True)

        with col_r:
            if net_iocs5:
                st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
                section_title("📡", f"NETWORK IOCs ({len(net_iocs5)})")
                for ioc in net_iocs5:
                    st.markdown(f'<div style="font-family:var(--font-mono);font-size:11px;color:var(--yellow);padding:4px 0;border-bottom:1px solid var(--border);">{_escape_html(str(ioc))}</div>', unsafe_allow_html=True)
                st.markdown('</div>', unsafe_allow_html=True)

            if recs5:
                st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
                section_title("⚡", f"RECOMMENDATIONS ({len(recs5)})")
                for i, rec in enumerate(recs5, 1):
                    st.markdown(f'<div style="display:flex;gap:12px;padding:8px 0;border-bottom:1px solid var(--border);"><span style="font-family:var(--font-mono);font-size:11px;color:var(--magenta);">{i:02d}</span><span style="font-size:12px;color:var(--text-secondary);line-height:1.5;">{_escape_html(str(rec))}</span></div>', unsafe_allow_html=True)
                st.markdown('</div>', unsafe_allow_html=True)

            if ui_behav5 and ui_behav5 != "—":
                st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
                section_title("🧭", "UI BEHAVIOUR")
                st.markdown(f'<div style="font-family:var(--font-body);font-size:13px;color:var(--text-secondary);line-height:1.7;white-space:pre-wrap;">{_escape_html(str(ui_behav5))}</div>', unsafe_allow_html=True)
                st.markdown('</div>', unsafe_allow_html=True)
    else:
        st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
        section_title("📄", "PHASE 5 — NOT YET RUN")
        st.markdown("""
        <div style="font-family:var(--font-body);font-size:13px;color:var(--text-secondary);line-height:1.7;">
        Phase 5 will generate a structured human-readable threat report from all pipeline outputs. Planned sections:<br><br>
        <span style="color:var(--cyan);">1.</span> Executive Summary<br>
        <span style="color:var(--cyan);">2.</span> Threat Classification &amp; Kill Chain<br>
        <span style="color:var(--cyan);">3.</span> Network IOCs<br>
        <span style="color:var(--cyan);">4.</span> Permission Abuse<br>
        <span style="color:var(--cyan);">5.</span> Dropper / Packer Chain<br>
        <span style="color:var(--cyan);">6.</span> UI Behaviour (Forms, Dialogs)<br>
        <span style="color:var(--cyan);">7.</span> Recommendations<br><br>
        <span style="color:var(--text-dim);">Until Phase 5 is implemented, use the Download Report button above for a Word document covering all completed phases.</span>
        </div>
        """, unsafe_allow_html=True)
        st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('</div>', unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════════
#  SCREENSHOTS
# ══════════════════════════════════════════════════════════════════════════════

elif "SCREENSHOTS" in page:
    pkg_str = _safe_get(SE, "package") or _safe_get(S, "package", default="—")

    st.markdown(f"""
    <div class="top-bar">
      <div style="display:flex;align-items:center;gap:16px;">
        <span class="brand">Dyn<span>Mal</span>Tool</span>
        <span style="font-family:var(--font-mono);font-size:11px;color:var(--cyan);">SCREENSHOTS · ALL PHASES</span>
        <span style="font-family:var(--font-mono);font-size:10px;color:var(--text-dim);">PKG · {_escape_html(str(pkg_str))}</span>
      </div>
      <span style="font-family:var(--font-mono);font-size:11px;color:var(--text-dim);">{len(screenshots)} image{"s" if len(screenshots)!=1 else ""}</span>
    </div>""", unsafe_allow_html=True)

    st.markdown('<div style="padding:20px 28px;">', unsafe_allow_html=True)

    if not screenshots:
        no_data(
            "No screenshots found in this session.\n\n"
            "Screenshots are captured when screenshot_tier < 3 (ADB screencap available).\n"
            "Check SCREENSHOT TIER in the sidebar — tier 3 = XML only, no PNGs."
        )
    else:
        # Grouping: main screenshots vs dropper subdirs
        main_shots = []
        dropper_shots: dict = {}

        for p in screenshots:
            # Dropper path: sessions/<s>/droppers/<pkg>/screenshots/<file>
            parts = p.parts
            try:
                d_idx = parts.index("droppers")
                pkg_name = parts[d_idx + 1]
                dropper_shots.setdefault(pkg_name, []).append(p)
            except (ValueError, IndexError):
                main_shots.append(p)

        # ── Filter / search bar ───────────────────────────────────────────────
        _fc, _sc = st.columns([3, 1])
        with _fc:
            ss_filter = st.text_input("Filter by filename", placeholder="e.g. intervention, dialog", label_visibility="collapsed", key="ss_filter")
        with _sc:
            ss_cols = st.selectbox("Grid width", [3, 4, 5, 6], index=1, label_visibility="collapsed", key="ss_cols")

        def _render_ss_group(paths: list, group_label: str):
            """Render a screenshot group as an image grid."""
            if ss_filter:
                paths = [p for p in paths if ss_filter.lower() in p.name.lower()]
            if not paths:
                no_data(f"No screenshots match filter in {group_label}")
                return

            st.markdown(f'<div style="font-family:var(--font-display);font-size:10px;letter-spacing:2px;color:var(--text-dim);margin:10px 0 8px;">{_escape_html(group_label)} — {len(paths)} image{"s" if len(paths)!=1 else ""}</div>', unsafe_allow_html=True)

            cols = st.columns(ss_cols)
            for i, img_path in enumerate(paths):
                with cols[i % ss_cols]:
                    try:
                        img_bytes = img_path.read_bytes()
                        b64 = base64.b64encode(img_bytes).decode()
                        fname = img_path.name
                        # Classify screenshot type by filename prefix
                        if "intervention" in fname:
                            border_color = "var(--orange)"
                        elif "dialog" in fname:
                            border_color = "var(--yellow)"
                        elif "dropper" in fname or "packer" in fname:
                            border_color = "var(--magenta)"
                        else:
                            border_color = "var(--border)"

                        st.markdown(
                            f'<div class="ss-card" style="border-color:{border_color};">'
                            f'<img src="data:image/png;base64,{b64}" style="width:100%;display:block;" />'
                            f'<div class="ss-label" title="{_escape_html(str(img_path))}">{_escape_html(fname)}</div>'
                            f'</div>',
                            unsafe_allow_html=True,
                        )
                        # Native download button per screenshot
                        st.download_button(
                            label="⬇",
                            data=img_bytes,
                            file_name=fname,
                            mime="image/png",
                            key=f"dl_ss_{i}_{hash(str(img_path))}",
                            use_container_width=True,
                        )
                    except OSError:
                        st.markdown(f'<div style="font-family:var(--font-mono);font-size:9px;color:var(--red);">⚠ {_escape_html(img_path.name)}</div>', unsafe_allow_html=True)
                    except Exception as e:
                        st.markdown(f'<div style="font-family:var(--font-mono);font-size:9px;color:var(--red);">Error: {_escape_html(str(e))}</div>', unsafe_allow_html=True)

        # Main screenshots
        if main_shots:
            st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
            section_title("📷", "MAIN APP SCREENSHOTS")
            _render_ss_group(main_shots, "Main app")
            st.markdown('</div>', unsafe_allow_html=True)
        else:
            st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
            section_title("📷", "MAIN APP SCREENSHOTS")
            no_data("No main-app screenshots found")
            st.markdown('</div>', unsafe_allow_html=True)

        # Dropper screenshots
        if dropper_shots:
            for pkg_name, d_paths in sorted(dropper_shots.items()):
                st.markdown('<div class="dmt-panel">', unsafe_allow_html=True)
                section_title("📦", f"DROPPER: {pkg_name}")
                _render_ss_group(d_paths, pkg_name)
                st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('</div>', unsafe_allow_html=True)


# ══════════════════════════════════════════════════════════════════════════════
#  LIVE LOGS
# ══════════════════════════════════════════════════════════════════════════════

elif "LOGS" in page:
    pkg_str = (_safe_get(SE, "package_name") or _safe_get(SE, "package") or
               _safe_get(S,  "package_name") or _safe_get(S,  "package") or "—")
    _pids   = _safe_get(SE, "all_pids", default=[]) or []
    pid_str = ", ".join(str(p) for p in _pids) if _pids else str(_safe_get(SE, "pid") or "—")

    st.markdown(f"""
    <div class="top-bar">
      <div style="display:flex;align-items:center;gap:16px;">
        <span class="brand">Dyn<span>Mal</span>Tool</span>
        <span style="font-family:var(--font-mono);font-size:11px;color:var(--cyan);">LOGCAT · PHASE_2</span>
        <span style="font-family:var(--font-mono);font-size:10px;color:var(--text-dim);">PKG · {_escape_html(str(pkg_str))}</span>
        <span style="font-family:var(--font-mono);font-size:10px;color:var(--text-dim);">PID · {_escape_html(pid_str)}</span>
      </div>
    </div>""", unsafe_allow_html=True)

    st.markdown('<div style="padding:20px 28px;">', unsafe_allow_html=True)

    log_lines = []

    if logcat_path and Path(logcat_path).exists():
        try:
            raw = Path(logcat_path).read_text(errors="replace").splitlines()
            pat = re.compile(r'^\S+\s+(\S+)\s+\d+\s+\d+\s+([A-Z])\s+\S+\s*:\s*(.*)')
            lvl_map = {"I":"INFO","W":"WARN","E":"ERROR","D":"INFO","V":"INFO","F":"ERROR"}
            for line in raw[-2000:]:
                try:
                    m = pat.match(line)
                    if m:
                        log_lines.append((m.group(1), lvl_map.get(m.group(2),"INFO"), m.group(3)))
                    else:
                        log_lines.append(("—","INFO",line))
                except Exception:
                    log_lines.append(("—","INFO",line))
        except OSError as e:
            st.warning(f"Could not read logcat file: {e}", icon="⚠️")
        except Exception as e:
            st.warning(f"Error parsing logcat: {e}", icon="⚠️")
    else:
        st.info("No `logcat.txt` found in session dir. Run Phase 2 to capture logs.", icon="ℹ️")

    if log_lines:
        col_f, col_l = st.columns([3, 2])
        with col_f:
            filter_text  = st.text_input("Log filter", placeholder="Filter (regex)…", label_visibility="collapsed")
        with col_l:
            level_filter = st.selectbox("Log level", ["ALL","INFO","WARN","ERROR"], label_visibility="collapsed")

        displayed = log_lines
        if level_filter != "ALL":
            displayed = [(ts,lv,msg) for ts,lv,msg in displayed if lv == level_filter]
        if filter_text:
            try:
                pat2 = re.compile(filter_text, re.IGNORECASE)
                displayed = [(ts,lv,msg) for ts,lv,msg in displayed if pat2.search(msg)]
            except re.error:
                st.warning("Invalid regex — showing all lines.", icon="⚠️")
            except Exception:
                pass

        st.markdown(render_log_lines(displayed), unsafe_allow_html=True)
        st.markdown(f"""
        <div style="display:flex;align-items:center;justify-content:space-between;margin-top:12px;padding:8px 14px;
            background:var(--bg-panel);border:1px solid var(--border);border-radius:4px;
            font-family:var(--font-mono);font-size:10px;color:var(--text-dim);">
          <div style="display:flex;gap:20px;">
            <span>FILE: <span style="color:var(--text-secondary);">{_escape_html(str(logcat_path))}</span></span>
            <span>LINES: <span style="color:var(--text-secondary);">{len(displayed)} / {len(log_lines)}</span></span>
          </div>
          <span>PID: <span style="color:var(--cyan);">{_escape_html(pid_str)}</span></span>
        </div>""", unsafe_allow_html=True)
    elif logcat_path:
        no_data("Log file exists but is empty")
    else:
        no_data("No log data available — run Phase 2 first")

    st.markdown('</div>', unsafe_allow_html=True)